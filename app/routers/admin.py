import json
import secrets
from datetime import date, datetime, timedelta

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, Request, UploadFile
from fastapi.responses import FileResponse, RedirectResponse, Response
from sqlalchemy.orm import Session

from ..auth import hash_password, require_roles, require_user
from ..config import BACKUP_DIR
from ..database import get_db, now_local, today_local
from ..models import (ROLES, ST, OPEN_STATUSES, ApiKey, AuditLog, Company, Customer, Die, NotificationChannel, Problem,
                      ProdLine, RepairRequest, RootCause, User)
from ..services import backup, images, import_data, notify
from ..services.kpi import pm_due
from ..web import render
from .requests import audit

from ..auth import get_current_user as _gcu

NO_COMPANY_OK = ("/admin/companies", "/admin/backup", "/admin/system")


def _need_company(request: Request, user=Depends(_gcu)):
    """A brand-new server (first deploy / disaster recovery) has no company yet: only the pages that can create
    one or restore a backup are usable; everything else sends the superadmin to Companies instead of crashing."""
    if user and getattr(request.state, "company", None) is None and not request.url.path.startswith(NO_COMPANY_OK):
        raise HTTPException(status_code=307, headers={"Location": "/admin/companies?msg=ยังไม่มีบริษัท — สร้างบริษัทใหม่ หรือกู้คืนจาก Backup ที่หน้า Backup"})


router = APIRouter(prefix="/admin", dependencies=[Depends(_need_company)])
admin_only = require_roles("admin")
superadmin_only = require_roles("superadmin")


@router.get("")
def admin_home(request: Request, user: User = Depends(admin_only), db: Session = Depends(get_db)):
    company = request.state.company
    stats = images.storage_stats(db, company.id if company else None)
    counts = {"users": db.query(User).filter(User.company_id == company.id).count() if company else 0,
              "dies": db.query(Die).filter(Die.company_id == company.id).count() if company else 0,
              "requests": db.query(RepairRequest).filter(RepairRequest.company_id == company.id).count() if company else 0,
              "channels": db.query(NotificationChannel).filter(NotificationChannel.company_id == company.id).count() if company else 0}
    from ..services.onboarding import readiness
    return render(request, "admin/index.html", stats=stats, counts=counts, ready=readiness(db, company),
                  backups=backup.list_backups(company.code)[:5] if company else [], db=db)


# ---------------------------------------------------------------- companies
@router.get("/companies")
def companies(request: Request, user: User = Depends(superadmin_only), db: Session = Depends(get_db), msg: str = ""):
    return render(request, "admin/companies.html", rows=db.query(Company).order_by(Company.id).all(), msg=msg, db=db)


@router.post("/companies")
def company_save(request: Request, user: User = Depends(superadmin_only), db: Session = Depends(get_db), cid: int = Form(0),
                 code: str = Form(...), name: str = Form(...), dr_prefix: str = Form("DR"), dr_style: str = Form("running"),
                 logo_text: str = Form("AAPICO"), is_active: str = Form("1"), masters: str = Form("default"),
                 admin_username: str = Form(""), admin_name: str = Form(""), admin_password: str = Form("")):
    code = code.strip().upper()
    clash = db.query(Company).filter(Company.code == code, Company.id != (cid or 0)).first()
    if clash:
        return RedirectResponse(f"/admin/companies?msg=รหัสบริษัท {code} มีอยู่แล้ว", 303)
    c = db.get(Company, cid) if cid else Company()
    c.code, c.name, c.dr_prefix, c.dr_style, c.logo_text = code, name.strip(), dr_prefix.strip(), dr_style, logo_text.strip()
    c.is_active = is_active == "1"
    msg = "บันทึกแล้ว"
    if not cid:
        db.add(c)
        db.flush()
        if masters == "default":
            seed_company_masters(db, c)
        if admin_username.strip() and admin_password:
            un = admin_username.strip().lower()
            if db.query(User).filter(User.username == un).first():
                db.rollback()
                return RedirectResponse(f"/admin/companies?msg=username {un} มีอยู่แล้ว", 303)
            db.add(User(company_id=c.id, username=un, full_name=admin_name.strip() or f"Admin {code}", role="admin",
                        password_hash=hash_password(admin_password), must_change_pw=True))
        msg = f"สร้างบริษัท {code} แล้ว — สลับไปบริษัท {code} (แถบด้านบน) แล้วไปที่ Admin › 🚀 เริ่มต้นบริษัท"
        audit(db, request, "company_create", "company", code, name)
    db.commit()
    if not cid:          # switch the superadmin straight into the new company's start-up page
        from ..auth import COOKIE, cookie_kwargs, make_session_cookie
        resp = RedirectResponse(f"/admin/setup?msg={msg}", 303)
        resp.set_cookie(COOKIE, make_session_cookie(user.id, c.id), **cookie_kwargs(request))
        return resp
    return RedirectResponse(f"/admin/companies?msg={msg}", 303)


def seed_company_masters(db: Session, c: Company):
    """Default problem / root-cause lists taken from the AH Power Apps masters."""
    problems = ["ครีบขอบ", "ครีบรู", "ชิ้นงานเป็นรอยขูด", "Part เป็นรอย", "Part แตก-ร้าว", "Part ฉีก", "Gap NG", "ขนาดรู NG", "ตำแหน่ง รู NG",
                "Trim Line NG", "คมตัด-แตกร้าว", "Insert แตก", "น็อต Insert ขาด", "PAD แตก", "น๊อต PAD ขาด", "ไกด์ PAD หัก", "ไกด์หัก", "ปลอกไกด์แตก",
                "แก๊สสปริงแตก", "ลิปเตอร์ไม่เด้ง", "แนวเชื่อมแตกร้าว", "Die ตก", "PM", "อื่นๆ"]
    causes = ["Punchสึก+Buttonสึก", "Cutting Edge - คมตัดแตก", "Part Burr - ครีบ", "Data NG - Gap NG", "Locater หัก", "Part เป็นรอย",
              "น็อต PAD ขาด", "Spring Broken - สปริงหัก", "Button Die Broken - บันทอม แตก บิ่น", "Punch Broken - พั้นช์หัก", "แก๊สสปริงแตก",
              "Insert แตก", "Guide Broken - ไกด์ แตก", "ไกด์หลวม", "PAD แตก", "Bolt Replaced", "Reuse STD", "Accident Die - อุบัติเหตุแม่พิมพ์ตก",
              "เศษ SCRAP ตัน", "Part Scratched", "Part แตกร้าว", "อื่นๆ"]
    for i, p in enumerate(problems):
        db.add(Problem(company_id=c.id, name_th=p, name_en=p, sort=i))
    for i, p in enumerate(causes):
        db.add(RootCause(company_id=c.id, name_th=p, name_en=p, sort=i))
    db.flush()


# ---------------------------------------------------------------- settings
@router.get("/settings")
def settings_page(request: Request, user: User = Depends(admin_only), db: Session = Depends(get_db), msg: str = ""):
    company = request.state.company
    s = company.settings
    return render(request, "admin/settings.html", s=s, shifts_json=json.dumps(s["shifts"], ensure_ascii=False, indent=1),
                  keys=db.query(ApiKey).filter(ApiKey.company_id == company.id).all(), msg=msg, db=db)


@router.post("/settings")
async def settings_save(request: Request, user: User = Depends(admin_only), db: Session = Depends(get_db)):
    company = db.get(Company, request.state.company.id)
    form = await request.form()
    try:
        shifts = json.loads(form.get("shifts_json") or "{}")
        import re as _re
        for name, sh in shifts.items():
            for k in ("start", "end"):
                if not _re.match(r"^\d{1,2}:\d{2}$", str(sh.get(k, ""))):
                    raise ValueError(f"{name}.{k}")
            for b in sh.get("breaks", []):
                if len(b) != 2 or not all(_re.match(r"^\d{1,2}:\d{2}$", str(x)) for x in b):
                    raise ValueError(f"{name}.breaks")
    except (json.JSONDecodeError, ValueError, AttributeError, TypeError) as e:
        return RedirectResponse(f"/admin/settings?msg=⚠ รูปแบบเวลากะ/เวลาพักไม่ถูกต้อง ({e}) — ใช้ HH:MM", 303)

    def num(key, default, lo=0, hi=100000):
        try:
            return max(lo, min(hi, int(float(form.get(key) or default))))
        except (TypeError, ValueError):
            return default
    company.set_settings({
        "pm_skip_foreman": bool(form.get("pm_skip_foreman")),
        "emergency_required_day": form.get("emergency_required_day") or "<1",
        "spare_app_url": (form.get("spare_app_url") or "").strip().rstrip("/"),
        "spare_api_key": (form.get("spare_api_key") or "").strip(),
        "image_archive_after_days": num("image_archive_after_days", 180, 7),
        "storage_warn_mb": num("storage_warn_mb", 400, 10),
        "approval_reminder_hours": num("approval_reminder_hours", 0),
        "pm_auto_create_days": num("pm_auto_create_days", 0, 0, 60),
        "pm_due_stroke_pct": num("pm_due_stroke_pct", 90, 50, 100),
        "pm_require_checklist": bool(form.get("pm_require_checklist")),
        "tv_working_hours": (form.get("tv_working_hours") or "").strip(),
        "shifts": shifts,
    })
    company.dr_prefix = (form.get("dr_prefix") or company.dr_prefix).strip()
    company.dr_style = form.get("dr_style") or company.dr_style
    company.logo_text = (form.get("logo_text") or company.logo_text).strip()
    company.name = (form.get("name") or company.name).strip()
    audit(db, request, "settings", "company", company.id)
    db.commit()
    return RedirectResponse("/admin/settings?msg=บันทึกแล้ว", 303)


@router.post("/apikeys")
def apikey_create(request: Request, user: User = Depends(admin_only), db: Session = Depends(get_db), name: str = Form("Spare Part App")):
    k = ApiKey(company_id=request.state.company.id, name=name.strip(), key="dms_" + secrets.token_urlsafe(24))
    db.add(k)
    db.commit()
    return RedirectResponse("/admin/settings", 303)


@router.post("/apikeys/{kid}/delete")
def apikey_delete(request: Request, kid: int, user: User = Depends(admin_only), db: Session = Depends(get_db)):
    k = db.get(ApiKey, kid)
    if k and k.company_id == request.state.company.id:
        db.delete(k)
        db.commit()
    return RedirectResponse("/admin/settings", 303)


# ---------------------------------------------------------------- users
@router.get("/users")
def users(request: Request, user: User = Depends(admin_only), db: Session = Depends(get_db), msg: str = ""):
    company = request.state.company
    rows = db.query(User).filter(User.company_id == company.id).order_by(User.role, User.full_name).all()
    lines = [l.name for l in db.query(ProdLine).filter_by(company_id=company.id).order_by(ProdLine.name)]
    return render(request, "admin/users.html", rows=rows, lines=lines, msg=msg, db=db)


@router.post("/users")
def user_save(request: Request, user: User = Depends(admin_only), db: Session = Depends(get_db), uid: int = Form(0),
              username: str = Form(...), full_name: str = Form(...), role: str = Form("production"), password: str = Form(""),
              emp_code: str = Form(""), email: str = Form(""), phone: str = Form(""), position: str = Form(""),
              department: str = Form(""), line: str = Form(""), shift: str = Form(""), skills: str = Form(""), is_active: str = Form("1"),
              must_change_pw: str = Form("")):
    company = request.state.company
    if role not in ROLES or (role == "superadmin" and not user.is_superadmin):
        raise HTTPException(400, "Invalid role")
    u = db.get(User, uid) if uid else None
    if u and u.company_id != company.id and not user.is_superadmin:
        raise HTTPException(403)
    username = username.strip().lower()
    clash = db.query(User).filter(User.username == username, User.id != (uid or 0)).first()
    if clash:
        return RedirectResponse(f"/admin/users?msg=ชื่อผู้ใช้ {username} มีอยู่แล้ว", 303)
    if not u:
        if not password:
            return RedirectResponse("/admin/users?msg=กรุณากำหนดรหัสผ่าน", 303)
        u = User(company_id=company.id, username=username, password_hash=hash_password(password))
        db.add(u)
    elif password:
        u.password_hash = hash_password(password)
    if password:
        u.must_change_pw = bool(must_change_pw)
    u.username, u.full_name, u.role = username, full_name.strip(), role
    u.emp_code, u.email, u.phone, u.position = emp_code.strip(), email.strip(), phone.strip(), position.strip()
    u.department, u.line, u.shift, u.skills = department.strip(), line.strip(), shift.strip(), skills.strip()
    u.is_active = is_active == "1"
    audit(db, request, "user_save", "user", u.id or 0, username)
    db.commit()
    return RedirectResponse("/admin/users?msg=บันทึกแล้ว", 303)


# ---------------------------------------------------------------- masters
@router.get("/masters")
def masters_page(request: Request, user: User = Depends(require_roles("admin", "die_leader", "manager")), db: Session = Depends(get_db), msg: str = ""):
    cid = request.state.company.id
    return render(request, "admin/masters.html",
                  customers=db.query(Customer).filter_by(company_id=cid).order_by(Customer.name).all(),
                  lines=db.query(ProdLine).filter_by(company_id=cid).order_by(ProdLine.name).all(),
                  problems=db.query(Problem).filter_by(company_id=cid).order_by(Problem.sort, Problem.name_th).all(),
                  causes=db.query(RootCause).filter_by(company_id=cid).order_by(RootCause.sort, RootCause.name_th).all(), msg=msg, db=db)


@router.post("/masters/{kind}")
def master_save(request: Request, kind: str, user: User = Depends(require_roles("admin", "die_leader", "manager")), db: Session = Depends(get_db),
                mid: int = Form(0), name: str = Form(...), name_en: str = Form(""), sort: int = Form(0), is_active: str = Form("1"), delete: str = Form("")):
    cid = request.state.company.id
    model = {"customer": Customer, "line": ProdLine, "problem": Problem, "cause": RootCause}.get(kind)
    if not model:
        raise HTTPException(404)
    obj = db.get(model, mid) if mid else model(company_id=cid)
    if obj is None or obj.company_id != cid:
        raise HTTPException(404)
    if delete and mid:
        used = 0
        if kind == "problem":
            used = db.query(RepairRequest).filter(RepairRequest.problem_id == obj.id).count()
        elif kind == "cause":
            used = db.query(RepairRequest).filter(RepairRequest.root_cause_id == obj.id).count()
        if used:          # referenced by tickets: keep the history, just hide it from the dropdowns
            obj.is_active = False
            db.commit()
            return RedirectResponse(f"/admin/masters?msg=⚠ รายการนี้ถูกใช้ใน {used} ใบงาน จึงเปลี่ยนเป็น 'ไม่ใช้งาน' แทนการลบ", 303)
        db.delete(obj)
        db.commit()
        return RedirectResponse("/admin/masters?msg=ลบแล้ว", 303)
    if kind in ("customer", "line"):
        obj.name = name.strip()
    else:
        obj.name_th, obj.name_en, obj.sort = name.strip(), name_en.strip() or name.strip(), sort
    obj.is_active = is_active == "1"
    if not mid:
        db.add(obj)
    db.commit()
    return RedirectResponse("/admin/masters?msg=บันทึกแล้ว", 303)


@router.post("/masters/clean/rename")
def master_rename(request: Request, user: User = Depends(require_roles("admin", "die_leader", "manager")), db: Session = Depends(get_db),
                  field: str = Form(...), old: str = Form(...), new: str = Form(...)):
    """Data cleaning: rename a value everywhere (requests + dies), e.g. 'LIne3' -> 'Line3'."""
    cid = request.state.company.id
    if field not in ("line", "customer", "die_process", "model", "problem_text", "root_cause_text"):
        raise HTTPException(400)
    n = db.query(RepairRequest).filter(RepairRequest.company_id == cid, getattr(RepairRequest, field) == old).update({field: new.strip()})
    if field in ("line", "customer", "die_process", "model"):
        n += db.query(Die).filter(Die.company_id == cid, getattr(Die, field) == old).update({field: new.strip()})
    audit(db, request, "rename", field, old, f"-> {new} ({n} rows)")
    db.commit()
    return RedirectResponse(f"/admin/masters?msg=เปลี่ยน {old} เป็น {new} แล้ว {n} รายการ", 303)


# ---------------------------------------------------------------- notification channels
@router.get("/channels")
def channels(request: Request, user: User = Depends(admin_only), db: Session = Depends(get_db), msg: str = ""):
    company = request.state.company
    rows = db.query(NotificationChannel).filter_by(company_id=company.id).all()
    approvers = (db.query(User).filter(User.company_id == company.id, User.is_active == True,
                                       User.role.in_(["production", "foreman", "die_leader", "technician", "qc", "manager"]))
                 .order_by(User.role, User.full_name).all())
    no_email = [u for u in approvers if not (u.email or "").strip() or "@" not in (u.email or "")]
    shared = {}
    for u in approvers:
        if u.email:
            shared.setdefault(u.email.strip().lower(), []).append(u.full_name)
    shared = {k: v for k, v in shared.items() if len(v) > 1}
    return render(request, "admin/channels.html", rows=rows, events=notify.EVENTS, msg=msg, pe=notify.email_config(company, db),
                  s=company.settings, no_email=no_email, shared=shared, n_approvers=len(approvers), db=db)


@router.post("/channels/personal-email")
def personal_email_save(request: Request, user: User = Depends(admin_only), db: Session = Depends(get_db),
                        enabled: str = Form(""), api_key: str = Form(""), sender_email: str = Form(""), sender_name: str = Form(""),
                        test_to: str = Form("")):
    c = db.get(Company, request.state.company.id)
    data = {"personal_email": bool(enabled), "brevo_sender_email": sender_email.strip(), "brevo_sender_name": sender_name.strip()}
    if api_key.strip() and not api_key.startswith("••"):
        data["brevo_api_key"] = api_key.strip()
    c.set_settings(data)
    audit(db, request, "personal_email_settings", "company", c.id)
    db.commit()
    msg = "บันทึกการตั้งค่าอีเมลรายบุคคลแล้ว"
    if test_to.strip():
        cfg = notify.email_config(c, db)
        if not cfg["api_key"] or not cfg["sender_email"]:
            msg += " · ยังส่งทดสอบไม่ได้: ต้องมี API key และอีเมลผู้ส่ง"
        else:
            import httpx
            try:
                r = httpx.post("https://api.brevo.com/v3/smtp/email", headers={"api-key": cfg["api_key"]}, timeout=15,
                               json={"sender": {"email": cfg["sender_email"], "name": cfg["sender_name"]}, "to": [{"email": test_to.strip()}],
                                     "subject": f"[{c.code}] ทดสอบอีเมลแจ้งผู้อนุมัติ", "htmlContent": "<p>ถ้าได้รับอีเมลนี้ แปลว่าตั้งค่าถูกต้อง ✅</p>"})
                msg += f" · ส่งทดสอบไป {test_to} ผล {r.status_code}" + ("" if r.status_code < 300 else f" {r.text[:120]}")
            except Exception as e:
                msg += f" · ส่งทดสอบไม่สำเร็จ: {e}"
    return RedirectResponse(f"/admin/channels?msg={msg}", 303)


@router.post("/channels")
async def channel_save(request: Request, user: User = Depends(admin_only), db: Session = Depends(get_db)):
    form = await request.form()
    cid = request.state.company.id
    chid = int(form.get("chid") or 0)
    ch = db.get(NotificationChannel, chid) if chid else NotificationChannel(company_id=cid)
    if ch.company_id != cid:
        raise HTTPException(403)
    if form.get("delete") and chid:
        db.delete(ch)
        db.commit()
        return RedirectResponse("/admin/channels", 303)
    ch.name, ch.type = form.get("name", "").strip(), form.get("type", "webhook")
    cfg = {k[4:]: v.strip() for k, v in form.items() if k.startswith("cfg_") and isinstance(v, str)}
    ch.config_json = json.dumps(cfg)
    ch.events_json = json.dumps(form.getlist("events"))
    ch.is_active = bool(form.get("is_active"))
    if not chid:
        db.add(ch)
    db.commit()
    if form.get("test"):
        try:
            res = notify.test_channel(ch.type, ch.config)
            return RedirectResponse(f"/admin/channels?msg=ส่งทดสอบแล้ว ({res})", 303)
        except Exception as e:
            return RedirectResponse(f"/admin/channels?msg=ส่งไม่สำเร็จ: {e}", 303)
    return RedirectResponse("/admin/channels?msg=บันทึกแล้ว", 303)


# ---------------------------------------------------------------- storage / images
@router.get("/storage")
def storage(request: Request, user: User = Depends(admin_only), db: Session = Depends(get_db), msg: str = ""):
    company = request.state.company
    stats = images.storage_stats(db, company.id)
    from sqlalchemy import func
    from ..models import Attachment
    by_month = (db.query(func.substr(Attachment.path, len(company.code) + 2, 7).label("ym"), func.count(Attachment.id),
                         func.coalesce(func.sum(Attachment.size_bytes), 0))
                .filter(Attachment.company_id == company.id).group_by("ym").order_by("ym").all())
    archivable = images.storage_stats(db, company.id)
    return render(request, "admin/storage.html", stats=stats, by_month=by_month, msg=msg,
                  archive_days=company.settings.get("image_archive_after_days", 180), db=db)


@router.post("/storage/archive")
def storage_archive(request: Request, user: User = Depends(admin_only), db: Session = Depends(get_db), days: int = Form(180)):
    n = images.archive_old(db, request.state.company.id, max(30, days))
    audit(db, request, "archive_images", "attachment", n)
    db.commit()
    return RedirectResponse(f"/admin/storage?msg=ลบรูปต้นฉบับของงานที่ปิดแล้วเกิน {days} วัน จำนวน {n} รูป (เก็บ thumbnail ไว้)", 303)


@router.post("/storage/export")
def storage_export(request: Request, user: User = Depends(admin_only), db: Session = Depends(get_db)):
    company = request.state.company
    name = f"images_{company.code}_{now_local():%Y%m%d_%H%M}.zip"
    n = images.export_images_zip(db, company.id, BACKUP_DIR / name)
    return RedirectResponse(f"/admin/backup?msg=รวมรูป {n} ไฟล์ไว้ใน {name}", 303)


# ---------------------------------------------------------------- backup
def _backup_code(request, user, scope="company"):
    if request.state.company is None:
        return "ALL"
    return "ALL" if (scope == "all" and user.is_superadmin) else request.state.company.code


def _visible_backup(request, user, name) -> str:
    """Returns the company code folder the file belongs to, or raises 404 if this admin may not see it."""
    if "/" in name or "\\" in name or ".." in name or not name.endswith(".zip"):
        raise HTTPException(404)
    codes = ([request.state.company.code] if request.state.company else []) + (["ALL"] if user.is_superadmin else [])
    for code in codes:
        if any(b["name"] == name for b in backup.list_backups(code)):
            return code
    raise HTTPException(404)


@router.get("/backup")
def backup_page(request: Request, user: User = Depends(admin_only), db: Session = Depends(get_db), msg: str = ""):
    company = request.state.company
    code = company.code if company else "ALL"
    items = backup.list_backups(code) if company else []
    if user.is_superadmin:
        items += [b for b in backup.list_backups("ALL") if b["name"] not in {i["name"] for i in items}]
    return render(request, "admin/backup.html", backups=items, s=company.settings if company else dict(Company.DEFAULT_SETTINGS),
                  s3=images.using_s3(), age=backup.last_backup_age_hours(code), no_company=company is None, db=db)


@router.post("/backup")
def backup_make(request: Request, user: User = Depends(admin_only), db: Session = Depends(get_db), scope: str = Form("company"),
                include_images: str = Form(""), send: str = Form("")):
    company = request.state.company
    cid = None if (company is None or (scope == "all" and user.is_superadmin)) else company.id
    code = _backup_code(request, user, scope)
    name = backup.make_backup(db, cid, bool(include_images), company_code=code)
    audit(db, request, "backup", "db", name)
    db.commit()
    sent = 0
    if send and company:
        blob = backup.get_backup_bytes(name, code)
        sent = notify.emit(db, company, "backup", None, title=f"[{company.code}] Backup {name}",
                           body=f"ไฟล์ Backup {round(len(blob) / 1048576, 2)} MB — เก็บไว้ในที่ปลอดภัย", link="/admin/backup",
                           attachment=(name, blob, "application/zip"))
        db.commit()
    where = "เซิร์ฟเวอร์ + Bucket" if images.using_s3() else "เซิร์ฟเวอร์ (กรุณาดาวน์โหลดเก็บ)"
    return RedirectResponse(f"/admin/backup?msg=สร้าง {name} แล้ว — เก็บที่ {where}" + (f" · ส่งไป {sent} ช่องทาง" if send else ""), 303)


@router.get("/backup/download/{name}")
def backup_download(request: Request, name: str, user: User = Depends(admin_only)):
    code = _visible_backup(request, user, name)
    data = backup.get_backup_bytes(name, code)
    return Response(data, media_type="application/zip", headers={"Content-Disposition": f"attachment; filename={name}"})


@router.post("/backup/delete/{name}")
def backup_delete(request: Request, name: str, user: User = Depends(admin_only), db: Session = Depends(get_db)):
    code = _visible_backup(request, user, name)
    backup.delete_backup(name, code)
    audit(db, request, "backup_delete", "db", name)
    db.commit()
    return RedirectResponse("/admin/backup?msg=ลบแล้ว", 303)


@router.post("/backup/settings")
def backup_settings(request: Request, user: User = Depends(admin_only), db: Session = Depends(get_db), auto_backup: str = Form(""),
                    auto_backup_hour: int = Form(2), backup_keep: int = Form(14), daily_report: str = Form(""),
                    daily_report_hour: int = Form(7)):
    if request.state.company is None:
        return RedirectResponse("/admin/companies?msg=สร้างบริษัทก่อนตั้งค่า Backup อัตโนมัติ", 303)
    c = db.get(Company, request.state.company.id)
    c.set_settings({"auto_backup": bool(auto_backup), "auto_backup_hour": max(0, min(23, auto_backup_hour)),
                    "backup_keep": max(1, min(90, backup_keep)), "daily_report": bool(daily_report),
                    "daily_report_hour": max(0, min(23, daily_report_hour))})
    audit(db, request, "backup_settings", "company", c.id)
    db.commit()
    return RedirectResponse("/admin/backup?msg=บันทึกการตั้งค่า Backup อัตโนมัติแล้ว", 303)


@router.post("/backup/restore")
async def backup_restore(request: Request, user: User = Depends(superadmin_only), db: Session = Depends(get_db),
                         file: UploadFile = File(...), replace: str = Form("")):
    data = await file.read()
    try:
        counts = backup.restore_backup(db, data, bool(replace))
    except Exception as e:
        db.rollback()
        return RedirectResponse(f"/admin/backup?msg=กู้คืนไม่สำเร็จ: {type(e).__name__}: {str(e)[:200]}", 303)
    audit(db, request, "restore", "db", file.filename, str(counts))
    db.commit()
    total = sum(v for k, v in counts.items() if k not in ("images", "skipped_conflicts"))
    # user ids may now belong to restored people -> end this session, log in again with the restored accounts
    from ..auth import COOKIE
    import urllib.parse as _u
    resp = RedirectResponse("/login?msg=" + _u.quote(f"กู้คืนแล้ว {total} แถว — กรุณาเข้าสู่ระบบใหม่ด้วยบัญชีจาก Backup · {counts}"), 303)
    resp.delete_cookie(COOKIE)
    return resp


# ---------------------------------------------------------------- import historic data / dies
@router.get("/import")
def import_page(request: Request, user: User = Depends(require_roles("admin", "die_leader", "manager")), db: Session = Depends(get_db), msg: str = ""):
    return render(request, "admin/import.html", msg=msg, db=db)


@router.get("/import/template.xlsx")
def import_template(user: User = Depends(require_user)):
    return Response(import_data.build_template(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": "attachment; filename=die_request_import_template.xlsx"})


@router.post("/import/requests")
async def import_requests(request: Request, user: User = Depends(require_roles("admin", "die_leader", "manager")), db: Session = Depends(get_db),
                          file: UploadFile = File(...)):
    data = await file.read()
    res = import_data.import_requests(db, request.state.company, file.filename, data, user)
    audit(db, request, "import_requests", "request", file.filename, str(res))
    db.commit()
    msg = f"นำเข้า {res['created']} รายการ ข้าม {res['skipped']} จาก {res['total_rows']} แถว" + (f" | ข้อผิดพลาด: {res['errors'][:5]}" if res["errors"] else "")
    return RedirectResponse(f"/admin/import?msg={msg}", 303)


@router.post("/import/dies")
async def import_dies(request: Request, user: User = Depends(require_roles("admin", "die_leader", "manager")), db: Session = Depends(get_db),
                      file: UploadFile = File(...)):
    data = await file.read()
    res = import_data.import_dies(db, request.state.company, file.filename, data)
    audit(db, request, "import_dies", "die", file.filename, str(res))
    db.commit()
    return RedirectResponse(f"/admin/import?msg=ทะเบียนแม่พิมพ์: เพิ่ม {res['created']} อัปเดต {res['updated']} จาก {res['total_rows']} แถว", 303)


@router.post("/import/close-stale")
def close_stale(request: Request, user: User = Depends(require_roles("admin", "manager")), db: Session = Depends(get_db), days: int = Form(60)):
    """Legacy jobs stuck at try-out / QC for a long time: close them in bulk (audit trail kept) so KPIs reflect reality."""
    from ..models import StatusHistory
    cutoff = now_local() - timedelta(days=max(7, days))
    rows = db.query(RepairRequest).filter(RepairRequest.company_id == request.state.company.id, RepairRequest.legacy_id != None,
                                          RepairRequest.status.in_([ST["REPAIR_DONE"], ST["WAIT_QC"]]), RepairRequest.created_at < cutoff).all()
    for r in rows:
        db.add(StatusHistory(request_id=r.id, from_status=r.status, to_status=ST["CLOSED"], user_id=user.id, user_name=user.full_name,
                             comment=f"ปิดอัตโนมัติ: งานนำเข้าค้างเกิน {days} วัน"))
        r.status, r.closed_at, r.closed_by_id = ST["CLOSED"], now_local(), user.id
        r.qc_result = r.qc_result or "OK"
    audit(db, request, "close_stale", "request", len(rows), f"> {days} days")
    db.commit()
    return RedirectResponse(f"/admin/import?msg=ปิดงานเก่าที่ค้าง Try-out/QC เกิน {days} วัน จำนวน {len(rows)} รายการ", 303)


# ---------------------------------------------------------------- audit
@router.get("/audit")
def audit_page(request: Request, user: User = Depends(admin_only), db: Session = Depends(get_db), page: int = 1,
               q: str = "", action: str = ""):
    from sqlalchemy import or_ as _or
    qry = db.query(AuditLog).filter(AuditLog.company_id == request.state.company.id)
    if action:
        qry = qry.filter(AuditLog.action == action)
    if q:
        like = f"%{q.strip()}%"
        qry = qry.filter(_or(AuditLog.user_name.ilike(like), AuditLog.entity_id.ilike(like), AuditLog.detail.ilike(like)))
    rows = qry.order_by(AuditLog.id.desc()).offset((page - 1) * 100).limit(100).all()
    actions = [a[0] for a in db.query(AuditLog.action).filter(AuditLog.company_id == request.state.company.id).distinct().order_by(AuditLog.action)]
    return render(request, "admin/audit.html", rows=rows, page=page, q=q, action=action, actions=actions, db=db)


# ---------------------------------------------------------------- scheduled checks (also runs automatically every hour)
def run_checks(db: Session, company: Company) -> dict:
    """Overdue jobs, PM due, storage warning -> notifications. Idempotent per day."""
    today = today_local()
    marker = f"checks_{today.isoformat()}"
    s = company.settings
    if s.get("_last_check") == marker:
        return {"skipped": True}
    overdue = db.query(RepairRequest).filter(RepairRequest.company_id == company.id, RepairRequest.is_deleted == False,
                                             RepairRequest.status.in_(OPEN_STATUSES),
                                             RepairRequest.target_date < today).all()
    if overdue:
        body = "\n".join(f"{r.dr_no} • {r.part_no} • {r.status} • กำหนด {r.target_date:%d-%m-%Y}" for r in overdue[:15])
        notify.emit(db, company, "overdue", None, title=f"[{company.code}] งานเกินกำหนด {len(overdue)} รายการ", body=body, link="/requests?view=open")
    pm = pm_due(db, company.id, 7)
    if pm["overdue"] or pm["due"]:
        body = "\n".join(f"{d.part_no} | {d.die_process} • " + (f"{d.pm_next_date:%d-%m-%Y}" if d.pm_next_date else (f"stroke {d.stroke_pct}%" if d.pm_interval_strokes else "ยังไม่เคย PM"))
                          for d in (pm["overdue"] + pm["due"])[:15])
        notify.emit(db, company, "pm_due", None, title=f"[{company.code}] PM เกินกำหนด {len(pm['overdue'])} / ถึงกำหนดใน 7 วัน {len(pm['due'])}", body=body, link="/pm")
    st = images.storage_stats(db, company.id)
    if st["mb"] >= s.get("storage_warn_mb", 400):
        notify.emit(db, company, "storage_warning", None, title=f"[{company.code}] พื้นที่รูปภาพ {st['mb']} MB",
                    body="กรุณา Backup รูปแล้วกด Archive ที่ Admin > Storage", link="/admin/storage")
    # personal reminders: tickets stuck waiting for the same person longer than N hours
    reminded = 0
    hrs = int(s.get("approval_reminder_hours", 8) or 0)
    if hrs > 0:
        waiting_event = {ST["WAIT_FOREMAN"]: "submit", ST["WAIT_LEADER"]: "foreman_approve", ST["REPAIR_DONE"]: "complete",
                         ST["WAIT_QC"]: "tryout_accept", ST["RETURNED"]: "tryout_return", ST["NG_RETURN"]: "qc_ng"}
        stuck = (db.query(RepairRequest).filter(RepairRequest.company_id == company.id, RepairRequest.is_deleted == False,
                                                RepairRequest.status.in_(list(waiting_event)),
                                                RepairRequest.updated_at < now_local() - timedelta(hours=hrs))
                 .order_by(RepairRequest.updated_at).limit(200).all())
        for r in stuck:
            waited = int((now_local() - r.updated_at).total_seconds() // 3600)
            try:
                notify.route_to_people(db, company, waiting_event[r.status], r,
                                       f"[{company.code}] ⏰ เตือน: ค้างที่คุณ {waited} ชม.", notify._req_summary(r), f"/requests/{r.id}",
                                       reminder=True)
                reminded += 1
            except Exception as e:
                print(f"[reminder] {r.dr_no}: {e}")
    keep_days = int(s.get("notification_keep_days", 90))
    from ..models import Notification
    old_n = (db.query(Notification).filter(Notification.company_id == company.id,
                                           Notification.at < now_local() - timedelta(days=keep_days))
             .delete(synchronize_session=False))
    c = db.get(Company, company.id)
    c.set_settings({"_last_check": marker})
    db.commit()
    return {"reminders": reminded, "notifications_cleaned": old_n, "overdue": len(overdue), "pm_overdue": len(pm["overdue"]), "pm_due": len(pm["due"]), "storage_mb": st["mb"]}


@router.post("/run-checks")
def run_checks_now(request: Request, user: User = Depends(admin_only), db: Session = Depends(get_db)):
    c = db.get(Company, request.state.company.id)
    c.set_settings({"_last_check": ""})
    db.commit()
    res = run_checks(db, c)
    return RedirectResponse(f"/admin?msg={res}", 303)


# ---------------------------------------------------------------- daily automatic jobs (called by the hourly scheduler)
def run_daily_jobs(db: Session, company: Company) -> dict:
    """Auto-backup + daily report, each at most once per day, after the configured local hour."""
    s = company.settings
    now = now_local()
    today = now.date().isoformat()
    done = {}
    c = db.get(Company, company.id)
    if s.get("auto_backup", True) and now.hour >= int(s.get("auto_backup_hour", 2)) and s.get("_last_auto_backup") != today:
        name = backup.make_backup(db, company.id, False, auto=True, company_code=company.code)
        pruned = backup.prune_auto_backups(company.code, int(s.get("backup_keep", 14)))
        try:
            blob = backup.get_backup_bytes(name, company.code)
            notify.emit(db, company, "backup", None, title=f"[{company.code}] Auto backup {name}",
                        body=f"Backup อัตโนมัติประจำวัน {round(len(blob) / 1048576, 2)} MB", link="/admin/backup",
                        attachment=(name, blob, "application/zip"))
        except Exception as e:
            print(f"[auto-backup] notify failed: {e}")
        c.set_settings({"_last_auto_backup": today})
        db.add(AuditLog(company_id=company.id, user_name="system", action="auto_backup", entity="db", entity_id=name,
                        detail=f"pruned {pruned}"))
        db.commit()
        done["backup"] = name
    if int(s.get("pm_auto_create_days") or 0) > 0 and now.hour >= 6 and s.get("_last_pm_auto") != today:
        from ..services import pm as pm_service
        made = pm_service.auto_create(db, c)
        c.set_settings({"_last_pm_auto": today})
        db.commit()
        if made:
            body = "\n".join(f"{r.dr_no} • {r.part_no} {r.die_process} • กำหนด {r.target_date:%d-%m-%Y}" for r in made[:20])
            notify.emit(db, company, "pm_due", None, title=f"[{company.code}] เปิดใบงาน PM อัตโนมัติ {len(made)} ใบ", body=body,
                        link="/requests?view=open&rtype=PM")
            for r in made[:10]:                 # personal e-mails for the first ones; the summary above lists all
                notify.emit(db, company, "submit", r)
            db.commit()
        done["pm_auto"] = len(made)
    if s.get("daily_report") and now.hour >= int(s.get("daily_report_hour", 7)) and s.get("_last_daily_report") != today:
        from ..services import export, kpi
        from ..models import RepairRequest
        from sqlalchemy.orm import selectinload
        rows = (db.query(RepairRequest).options(selectinload(RepairRequest.assignments))
                .filter(RepairRequest.company_id == company.id, RepairRequest.is_deleted == False,
                        RepairRequest.status.in_(OPEN_STATUSES)).order_by(RepairRequest.created_at).all())
        sm = kpi.summary(db, company.id, today_local().replace(day=1), today_local())
        body = (f"งานค้าง {len(rows)} · ด่วน {sum(1 for r in rows if r.is_emergency)} · เกินกำหนด {sum(1 for r in rows if r.is_overdue)}\n"
                f"เดือนนี้: แจ้ง {sm['total']} ปิด {sm['closed']} · MTTR {sm['mttr_text']} · QC NG {sm['ng_total']}")
        xlsx = export.to_xlsx(rows, "Open jobs")
        notify.emit(db, company, "daily_report", None, title=f"[{company.code}] รายงานประจำวัน {now:%d-%m-%Y}", body=body,
                    link="/requests?view=open",
                    attachment=(f"open_jobs_{company.code}_{now:%Y%m%d}.xlsx", xlsx,
                                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))
        c.set_settings({"_last_daily_report": today})
        db.commit()
        done["report"] = True
    return done


# ---------------------------------------------------------------- purge old jobs (start fresh, keep masters)
@router.get("/purge")
def purge_page(request: Request, user: User = Depends(admin_only), db: Session = Depends(get_db), scope: str = "legacy",
               types: list[str] = Query(default=["REPAIR", "PM", "ACCIDENT"]), status: str = "any", dfrom: str = "", dto: str = ""):
    from ..services import purge
    company = request.state.company
    pv = purge.preview(db, company.id, scope, types, status, dfrom, dto)
    placeholder = db.query(User).filter(User.company_id == company.id, User.password_hash == "!").count()
    nocode = db.query(Die).filter(Die.company_id == company.id, (Die.die_code == None) | (Die.die_code == "")).count()
    return render(request, "admin/purge.html", pv=pv, scope=scope, types=types, status=status, dfrom=dfrom, dto=dto,
                  scopes=purge.SCOPES, statuses=purge.STATUS_FILTERS, placeholder=placeholder, nocode=nocode, db=db)


@router.post("/purge")
async def purge_run(request: Request, user: User = Depends(admin_only), db: Session = Depends(get_db)):
    from ..services import purge
    company = request.state.company
    f = await request.form()
    if (f.get("confirm") or "").strip().upper() != company.code.upper():
        return RedirectResponse(f"/admin/purge?msg=ยกเลิก: พิมพ์รหัสบริษัท {company.code} ให้ถูกต้องเพื่อยืนยัน", 303)
    scope, status = f.get("scope", "legacy"), f.get("status", "any")
    types = f.getlist("types") or ["REPAIR", "PM", "ACCIDENT"]
    dfrom, dto = f.get("dfrom", ""), f.get("dto", "")
    # 1) mandatory safety backup (includes photos so the purge is fully reversible)
    bname = backup.make_backup(db, company.id, include_images=True, company_code=company.code)
    nd = (f.get("next_dr") or "").strip()
    res = purge.execute(db, company, scope, types, status, dfrom, dto,
                        delete_placeholder_users=bool(f.get("del_users")), delete_orphan_dies=bool(f.get("del_dies")),
                        reset_pm_dates=bool(f.get("reset_pm")), clear_notifications=bool(f.get("clear_notif")),
                        next_dr=int(nd) if nd.isdigit() else None)
    db.add(AuditLog(company_id=company.id, user_id=user.id, user_name=user.full_name, action="purge", entity="request",
                    entity_id=str(res["requests"]), detail=f"scope={scope} types={types} status={status} {dfrom}-{dto} backup={bname} result={res}"))
    db.commit()
    from .dashboard import _TV_CACHE
    _TV_CACHE.clear()
    return RedirectResponse(f"/admin/purge?msg=ลบใบงาน {res['requests']} รายการแล้ว · Backup ก่อนลบ: {bname} (กู้คืนได้ที่หน้า Backup) · {res}", 303)


# ---------------------------------------------------------------- system health
@router.get("/system")
def system_page(request: Request, user: User = Depends(admin_only), db: Session = Depends(get_db)):
    import platform
    import sys
    import time as _t
    from sqlalchemy import text
    from .. import migrate
    from ..auth import is_local_install
    from ..config import settings as cfg, DATA_DIR
    from ..database import engine, Base
    from ..models import RepairRequest, Attachment
    company = request.state.company
    t0 = _t.time()
    db.execute(text("SELECT 1"))
    latency = round((_t.time() - t0) * 1000, 1)
    dialect = engine.dialect.name
    try:
        if dialect == "postgresql":
            db_mb = round(db.execute(text("SELECT pg_database_size(current_database())")).scalar() / 1048576, 1)
        else:
            f = DATA_DIR / "die_management.db"
            db_mb = round(f.stat().st_size / 1048576, 1) if f.exists() else 0
    except Exception:
        db_mb = None
    counts = {}
    for t in Base.metadata.sorted_tables:
        try:
            counts[t.name] = db.execute(text(f"SELECT COUNT(*) FROM {t.name}")).scalar()
        except Exception:
            counts[t.name] = "?"
    stats = images.storage_stats(db, company.id)
    age = backup.last_backup_age_hours(company.code)
    demo = db.query(User).filter(User.username.in_(["ah_admin", "ah_manager", "ah_prod", "ah_foreman", "ah_leader", "ah_tech1", "ah_tech2", "ah_qc"]),
                                 User.is_active == True).count()
    on_render = bool(__import__("os").getenv("RENDER"))
    checks = []

    def chk(ok, title, fix):
        checks.append({"ok": ok, "title": title, "fix": fix})

    chk(dialect == "postgresql" or not on_render, f"ฐานข้อมูล: {dialect}",
        "บน Render ดิสก์ถูกล้างทุกครั้งที่หลับ/deploy → ต้องตั้ง DATABASE_URL เป็น Postgres (Supabase แนะนำ)")
    chk(images.using_s3() or not on_render, f"ที่เก็บรูป: {stats['backend']}",
        "บน host ฟรีรูปบนดิสก์จะหาย → ตั้ง S3_* (Cloudflare R2 / Supabase Storage)")
    chk(cfg.SECRET_FROM_ENV or not cfg.BASE_URL.startswith("https"), "SECRET_KEY (กุญแจลงชื่อ cookie) ไม่ใช่ค่าเริ่มต้น",
        "บน host ฟรีให้ตั้ง ENV SECRET_KEY เป็นข้อความสุ่มยาว ๆ (render.yaml สร้างให้อัตโนมัติ) — ตอนนี้ใช้กุญแจสุ่มในไฟล์ data/.secret ซึ่งจะหายเมื่อ host ล้างดิสก์ (ทุกคนต้อง login ใหม่)")
    chk(is_local_install() or cfg.BASE_URL.startswith("https://"), f"BASE_URL = {cfg.BASE_URL}",
        "ตั้ง BASE_URL เป็น https://<ชื่อแอป>.onrender.com เพื่อให้ลิงก์ในแจ้งเตือน/QR ถูกต้อง")
    chk(is_local_install() or demo == 0, f"บัญชีทดลอง (รหัส 1234) ที่ยังเปิดอยู่: {demo}", "กดปุ่ม 'ปิดบัญชีทดลอง' ด้านล่าง ก่อนเปิดให้ใช้งานจริง")
    chk(age is not None and age < 30, f"Backup ล่าสุด: {age if age is not None else '-'} ชม.ที่แล้ว", "เปิด Backup อัตโนมัติ (หน้า Backup) หรือกดสร้างเอง")
    chk(not stats["warn"], f"พื้นที่รูป {stats['mb']} / {stats['warn_mb']} MB", "Export รูปแล้ว Archive ที่หน้า รูปภาพ/พื้นที่")
    if dialect == "postgresql" and db_mb is not None:
        chk(db_mb < 400, f"ขนาดฐานข้อมูล {db_mb} MB (Supabase ฟรี 500 MB / Neon ฟรี 1 GB)", "ลบใบงานเก่า (หน้าล้างข้อมูล) หรืออัปเกรดแพ็กเกจ")
    chk(latency < 150, f"เวลาตอบสนองฐานข้อมูล {latency} ms", "ให้แอปและฐานข้อมูลอยู่ region เดียวกัน (เช่น Singapore)")
    chans = db.query(NotificationChannel).filter(NotificationChannel.company_id == company.id, NotificationChannel.is_active == True).all()
    backup_out = images.using_s3() or any("backup" in (ch.events or []) and ch.type in ("telegram", "discord", "email") for ch in chans)
    chk(backup_out or not on_render, "Backup อัตโนมัติถูกส่งออกนอกเซิร์ฟเวอร์ (Bucket หรือ Telegram/Discord/Email)",
        "ดิสก์ host ฟรีไม่ถาวร → ตั้ง S3_* หรือเพิ่มช่องทางแจ้งเตือน Telegram/Discord/Email แล้วติ๊กเหตุการณ์ 'Backup'")
    chk(len(chans) > 0, f"ช่องทางแจ้งเตือนที่เปิดใช้: {len(chans)}", "เพิ่ม LINE / Telegram / Discord / Brevo ที่หน้า ช่องทางแจ้งเตือน")
    chk(not cfg.BASE_URL_AUTO or is_local_install(), "BASE_URL ยืนยันจากการใช้งานจริงแล้ว",
        "เปิดเว็บผ่านลิงก์ https ของ host หนึ่งครั้ง ระบบจะจำ URL เอง หรือตั้ง ENV BASE_URL")
    info = {"เวอร์ชันโปรแกรม": migrate.APP_VERSION, "Python": sys.version.split()[0], "OS": platform.platform(),
            "ฐานข้อมูล": dialect, "ขนาดฐานข้อมูล (MB)": db_mb, "Host": "Render" if on_render else ("เครื่องในโรงงาน" if is_local_install() else "อื่น ๆ"),
            "เวลาเซิร์ฟเวอร์ (โรงงาน)": now_local().strftime("%d-%m-%Y %H:%M:%S"),
            "ใบงานบริษัทนี้": db.query(RepairRequest).filter_by(company_id=company.id).count(),
            "รูปบริษัทนี้": db.query(Attachment).filter_by(company_id=company.id).count()}
    return render(request, "admin/system.html", checks=checks, info=info, counts=counts, db=db)


@router.post("/system/disable-demo")
def disable_demo(request: Request, user: User = Depends(admin_only), db: Session = Depends(get_db)):
    names = ["ah_admin", "ah_manager", "ah_prod", "ah_foreman", "ah_leader", "ah_tech1", "ah_tech2", "ah_qc"]
    q = db.query(User).filter(User.username.in_(names), User.id != user.id, User.company_id == request.state.company.id)
    n = q.update({User.is_active: False}, synchronize_session=False)
    audit(db, request, "disable_demo", "user", n)
    db.commit()
    note = " (บัญชีที่คุณใช้อยู่ไม่ถูกปิด — เปลี่ยนรหัสผ่านที่โปรไฟล์)" if user.username in names else ""
    return RedirectResponse(f"/admin/system?msg=ปิดบัญชีทดลอง {n} บัญชีแล้ว{note}", 303)


@router.post("/users/force-password")
def force_password_all(request: Request, user: User = Depends(admin_only), db: Session = Depends(get_db)):
    """Everybody (except you) must set a new password at next login."""
    n = (db.query(User).filter(User.company_id == request.state.company.id, User.id != user.id, User.is_active == True)
         .update({User.must_change_pw: True}, synchronize_session=False))
    audit(db, request, "force_password_all", "user", n)
    db.commit()
    return RedirectResponse(f"/admin/users?msg=ผู้ใช้ {n} คนต้องตั้งรหัสใหม่ในการเข้าครั้งถัดไป", 303)


# ---------------------------------------------------------------- master data export / user import
USER_COLS = ["username", "full_name", "role", "password", "emp_code", "position", "department", "line", "shift", "email", "phone", "skills", "is_active"]


@router.get("/export/{what}.xlsx")
def export_master(request: Request, what: str, user: User = Depends(require_roles("admin", "manager", "die_leader")), db: Session = Depends(get_db)):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    import io
    cid = request.state.company.id
    wb = Workbook()
    ws = wb.active
    if what == "dies":
        ws.title = "Dies"
        heads = ["Company", "Customer", "Model", "PartNo", "PartName", "DieProcess", "DieCode", "Line", "Machine", "PM Interval Days", "PM Last Date", "PM Next Date", "Active"]
        ws.append(heads)
        code = request.state.company.code
        for d in db.query(Die).filter(Die.company_id == cid).order_by(Die.customer, Die.part_no, Die.die_process):
            ws.append([code, d.customer, d.model, d.part_no, d.part_name, d.die_process, d.die_code, d.line, d.machine_no,
                       d.pm_interval_days, d.pm_last_date.strftime("%d-%m-%Y") if d.pm_last_date else "",
                       d.pm_next_date.strftime("%d-%m-%Y") if d.pm_next_date else "", "Y" if d.is_active else "N"])
    elif what in ("users", "users_template"):
        ws.title = "Users"
        ws.append(USER_COLS)
        if what == "users":
            for u in db.query(User).filter(User.company_id == cid).order_by(User.role, User.full_name):
                ws.append([u.username, u.full_name, u.role, "", u.emp_code, u.position, u.department, u.line, u.shift, u.email, u.phone, u.skills, "Y" if u.is_active else "N"])
        else:
            ws.append(["somchai.p", "นาย สมชาย ใจดี", "technician", "Temp#2026", "10001234", "Technician", "Die Maintenance", "", "DAY", "", "08x-xxx-xxxx", "เชื่อม, เจียร", "Y"])
            info = wb.create_sheet("วิธีใช้")
            for line in ["role ใช้ค่า: " + ", ".join(r for r in ROLES if r != "superadmin"),
                         "password: ว่าง = ไม่เปลี่ยนรหัสเดิม (ผู้ใช้ใหม่ต้องใส่) — ผู้ใช้ใหม่จะถูกบังคับเปลี่ยนรหัสเมื่อเข้าครั้งแรก",
                         "username ซ้ำกับที่มีอยู่ = อัปเดตข้อมูลคนนั้น", "is_active: Y / N"]:
                info.append([line])
    else:
        raise HTTPException(404)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F5FBF")
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 18
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename={what}_{request.state.company.code}.xlsx"})


@router.post("/users/import")
async def users_import(request: Request, user: User = Depends(admin_only), db: Session = Depends(get_db), file: UploadFile = File(...)):
    from ..services.onboarding import upsert_users
    raw = import_data.read_rows(file.filename, await file.read())
    rows = [{str(k).strip().lower().split(" ")[0]: v for k, v in r.items() if k} for r in raw]
    cid = request.state.company.id
    lines = {l.name.lower() for l in db.query(ProdLine).filter_by(company_id=cid)}
    res = upsert_users(db, cid, rows, hash_password, lines)
    audit(db, request, "users_import", "user", file.filename, f"+{res['created']} ~{res['updated']}")
    db.commit()
    errors = res["errors"]
    msg = f"นำเข้าผู้ใช้: เพิ่ม {res['created']} · อัปเดต {res['updated']}" + (f" · ปัญหา: {'; '.join(errors[:5])}" if errors else "")
    return RedirectResponse(f"/admin/users?msg={msg}", 303)
