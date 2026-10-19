import io
from datetime import datetime, timedelta

from fastapi import APIRouter, Depends, Request
from fastapi.responses import Response
from sqlalchemy.orm import Session

from .. import perms
from ..database import get_db, now_local, today_local
from ..models import RepairRequest, User
from ..services import analytics
from ..web import render

router = APIRouter()


@router.get("/analytics")
def analytics_page(request: Request, user: User = Depends(perms.require_page("analytics")), db: Session = Depends(get_db),
                   dfrom: str = "", dto: str = "", line: str = "", customer: str = "", rtype: str = "", fmt: str = ""):
    company = request.state.company
    today = today_local()
    start = datetime.strptime(dfrom, "%Y-%m-%d").date() if dfrom else (today - timedelta(days=365)).replace(day=1)
    end = datetime.strptime(dto, "%Y-%m-%d").date() if dto else today
    a = analytics.analyse(db, company.id, start, end, line, customer, rtype, now=now_local(), user=user)
    if fmt == "xlsx":
        return _excel(a, company.code, start, end)
    lines = sorted({r[0] for r in db.query(RepairRequest.line).filter(RepairRequest.company_id == company.id).distinct() if r[0]})
    customers = sorted({r[0] for r in db.query(RepairRequest.customer).filter(RepairRequest.company_id == company.id).distinct() if r[0]})
    return render(request, "analytics.html", a=a, dfrom=start.isoformat(), dto=end.isoformat(), line=line, customer=customer,
                  rtype=rtype, lines=lines, customers=customers, db=db)


def _excel(a, code, start, end):
    from ..services.export import dmy_text
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    sheets = [
        ("สรุป", ["ตัวชี้วัด", "ค่า"], [[k, v] for k, v in a["head"].items() if not isinstance(v, (dict, list))] + [["", ""]] + [["ข้อค้นพบ", t] for _, t, _ in a["insights"]]),
        ("Pareto ปัญหา", ["ปัญหา", "จำนวน", "%", "สะสม %"], [[p["name"], p["n"], p["pct"], p["cum"]] for p in a["pareto_problem"]]),
        ("Pareto สาเหตุ", ["สาเหตุ", "จำนวน", "%", "สะสม %"], [[p["name"], p["n"], p["pct"], p["cum"]] for p in a["pareto_cause"]]),
        ("ปัญหา x สาเหตุ", ["ปัญหา \\ สาเหตุ"] + a["matrix"]["cols"], [[r] + c for r, c in zip(a["matrix"]["rows"], a["matrix"]["cells"])]),
        ("แม่พิมพ์เสี่ยง", ["Part No", "Process", "ลูกค้า", "Line", "ซ่อม (ครั้ง)", "MTBF (วัน)", "ซ้ำใน 7 วัน", "ส่งกลับ/NG", "MTTR (นาที)", "ต้นทุน", "ปัญหาหลัก", "Risk"],
         [[d["part_no"], d["process"], d["customer"], d["line"], d["repairs"], d["mtbf"], d["rep7"], d["ng"], d["mttr"], d["cost"], d["top_problem"], d["risk"]] for d in a["dies"]]),
        ("Lead time ต่อขั้น", ["เดือน"] + a["stage_names"], [[m] + a["stage_month"][m] for m in a["months"]]),
        ("ทันกำหนด", ["Required day", "เป้า (ชม.)", "จำนวน", "ทันกำหนด %", "ใช้จริง median (ชม.)"],
         [[s["req"], s["target_h"], s["n"], s["ontime"], s["actual_h"]] for s in a["sla_by_req"]]),
        ("ช่าง", ["ช่าง", "งาน", "ปิด", "ค้าง", "เวลาซ่อมเฉลี่ย/คน (นาที)", "รอรับงาน median (ชม.)", "ผ่านครั้งเดียว %"],
         [[t["name"], t["jobs"], t["closed"], t["open"], t["avg_net"], t["recv_h"], t["ftp"]] for t in a["tech"]]),
        ("Line x เดือน", ["Line"] + a["heat"]["cols"], [[r] + c for r, c in zip(a["heat"]["rows"], a["heat"]["cells"])]),
        ("ชั่วโมงที่แจ้ง", ["ชั่วโมง", "จำนวน"], [[f"{i:02d}:00", n] for i, n in enumerate(a["hours"])]),
        ("งานค้างตามอายุ", ["สถานะ"] + a["age_labels"], [[g] + v for g, v in a["aging"].items()]),
    ]
    first = True
    for name, head, rows in sheets:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name[:30]
        ws.append([dmy_text(h) for h in head])
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F5FBF")
        for r in rows:
            ws.append([("" if v is None else dmy_text(v)) for v in r])
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 16
        ws.column_dimensions["A"].width = 34
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename=die_analytics_{code}_{start:%d-%m-%Y}_{end:%d-%m-%Y}.xlsx"})
