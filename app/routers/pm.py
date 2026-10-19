from datetime import date, datetime, timedelta

import io

from fastapi import APIRouter, Depends, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import RedirectResponse, Response
from sqlalchemy import or_
from sqlalchemy.orm import Session

from .. import workflow
from .. import perms
from ..auth import require_user, require_roles
from ..database import get_db, now_local, today_local
from ..models import ST, Die, PMEvent, PMTemplate, PMTemplateItem, RepairRequest, StrokeLog, User, OPEN_STATUSES
from ..services import notify
from ..services import pm as pm_service
from ..web import render
from .requests import audit, next_dr_no

router = APIRouter()


LEADS = ("die_leader", "admin", "manager")


def _lines_customers(db, cid):
    lines = sorted({r[0] for r in db.query(Die.line).filter(Die.company_id == cid, or_(Die.pm_interval_days > 0, Die.pm_interval_strokes > 0)).distinct() if r[0]})
    custs = sorted({r[0] for r in db.query(Die.customer).filter(Die.company_id == cid, or_(Die.pm_interval_days > 0, Die.pm_interval_strokes > 0)).distinct() if r[0]})
    return lines, custs


def _date(s):
    s = (s or "").strip()
    if not s:
        return None
    for f in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    raise HTTPException(400, f"วันที่ไม่ถูกต้อง: {s} (ใช้ DD-MM-YYYY)")


@router.get("/pm")
def pm_page(request: Request, user: User = Depends(perms.require_page("pm")), db: Session = Depends(get_db), horizon: int = 14,
            line: str = "", customer: str = "", show: str = "due"):
    company = request.state.company
    groups = pm_service.due_groups(db, company.id, horizon, line, customer)
    open_pm = (db.query(RepairRequest).filter(RepairRequest.company_id == company.id, RepairRequest.request_type == "PM",
                                              RepairRequest.is_deleted == False, RepairRequest.status.in_(OPEN_STATUSES))
               .order_by(RepairRequest.target_date).all())
    open_die_ids = {r.die_id for r in open_pm}
    today = today_local()
    k = pm_service.kpis(db, company.id, today.replace(day=1), today)
    k12 = pm_service.kpis(db, company.id, (today - timedelta(days=365)), today)
    lines, custs = _lines_customers(db, company.id)
    n_plan = db.query(Die).filter(Die.company_id == company.id, Die.is_active == True,
                                  or_(Die.pm_interval_days > 0, Die.pm_interval_strokes > 0)).count()
    n_dies = db.query(Die).filter(Die.company_id == company.id, Die.is_active == True).count()
    n_sug = len(pm_service.suggestions(db, company.id, limit=200))
    return render(request, "pm.html", groups=groups, open_pm=open_pm, open_die_ids=open_die_ids, horizon=horizon, line=line,
                  customer=customer, show=show, k=k, k12=k12, lines=lines, customers=custs, n_plan=n_plan, n_dies=n_dies,
                  n_sug=n_sug, db=db)


@router.post("/pm/create/{die_id}")
def pm_create(request: Request, die_id: int, user: User = Depends(require_roles("die_leader", "foreman", "admin", "manager")),
              db: Session = Depends(get_db), target_date: str = Form(""), note: str = Form("")):
    company = request.state.company
    die = db.get(Die, die_id)
    if not die or die.company_id != company.id:
        raise HTTPException(404)
    if die.id in pm_service.open_pm_die_ids(db, company.id):
        return RedirectResponse("/pm?msg=แม่พิมพ์นี้มีใบงาน PM ที่ยังเปิดอยู่แล้ว", 303)
    req = pm_service.open_pm_job(db, company, die, user, _date(target_date), note)
    audit(db, request, "create_pm", "request", req.id, req.dr_no)
    db.commit()
    notify.emit(db, company, "submit", req, actor_id=user.id)
    db.commit()
    return RedirectResponse(f"/requests/{req.id}?created=1", 303)


@router.post("/pm/create-all")
def pm_create_all(request: Request, user: User = Depends(require_roles(*LEADS)), db: Session = Depends(get_db),
                  horizon: int = Form(14), line: str = Form(""), customer: str = Form("")):
    """One click: open a PM job for every overdue / due die (current filter) that has no open PM job yet."""
    company = request.state.company
    groups = pm_service.due_groups(db, company.id, horizon, line, customer)
    busy = pm_service.open_pm_die_ids(db, company.id)
    made = []
    for die in groups["overdue"] + groups["due"]:
        if die.id not in busy:
            made.append(pm_service.open_pm_job(db, company, die, user))
    audit(db, request, "create_pm_all", "request", len(made), ", ".join(r.dr_no for r in made[:50]))
    db.commit()
    if made:
        body = "\n".join(f"{r.dr_no} • {r.part_no} {r.die_process} • กำหนด {r.target_date:%d-%m-%Y}" for r in made[:20])
        notify.emit(db, company, "pm_due", None, title=f"[{company.code}] เปิดใบงาน PM {len(made)} ใบ", body=body, link="/requests?view=open&rtype=PM")
        db.commit()
    return RedirectResponse(f"/pm?msg=สร้างใบงาน PM แล้ว {len(made)} ใบ", 303)


@router.post("/pm/auto-run")
def pm_auto_run(request: Request, user: User = Depends(require_roles(*LEADS)), db: Session = Depends(get_db)):
    """Run the daily automatic PM-job creation now (same rule as the scheduler)."""
    from ..models import Company
    company = db.get(Company, request.state.company.id)
    if int(company.settings.get("pm_auto_create_days") or 0) <= 0:
        return RedirectResponse("/pm?msg=⚠ ยังไม่ได้เปิดใช้ — ตั้ง 'เปิดใบงาน PM อัตโนมัติล่วงหน้า' ที่ ผู้ดูแลระบบ › ตั้งค่าบริษัท", 303)
    made = pm_service.auto_create(db, company)
    audit(db, request, "pm_auto_run", "request", len(made), ", ".join(r.dr_no for r in made[:50]))
    db.commit()
    for r in made[:10]:
        notify.emit(db, company, "submit", r)
    db.commit()
    return RedirectResponse(f"/pm?msg=เปิดใบงาน PM อัตโนมัติ {len(made)} ใบ", 303)


@router.post("/pm/postpone/{die_id}")
def pm_postpone(request: Request, die_id: int, user: User = Depends(require_roles(*LEADS)), db: Session = Depends(get_db),
                new_date: str = Form(...), reason: str = Form(...)):
    company = request.state.company
    die = db.get(Die, die_id)
    if not die or die.company_id != company.id or not die.pm_interval_days:
        raise HTTPException(404)
    nd = _date(new_date)
    if not reason.strip():
        raise HTTPException(400, "กรุณาระบุเหตุผลการเลื่อน")
    old = pm_service.next_date(die)
    die.pm_next_date = nd
    db.add(PMEvent(company_id=company.id, die_id=die.id, kind="POSTPONE", from_date=old, to_date=nd, reason=reason.strip()[:240],
                   user_name=user.full_name))
    for r in db.query(RepairRequest).filter(RepairRequest.die_id == die.id, RepairRequest.request_type == "PM",
                                            RepairRequest.status.in_(OPEN_STATUSES)):
        r.target_date = nd                                   # the open PM job follows the new date
    audit(db, request, "pm_postpone", "die", die.id, f"{old} -> {nd}: {reason}")
    db.commit()
    return RedirectResponse(f"/pm?msg=เลื่อน PM ของ {die.part_no} เป็น {nd:%d-%m-%Y} แล้ว", 303)


# ---------------------------------------------------------------- calendar + annual plan
@router.get("/pm/calendar")
def pm_calendar(request: Request, user: User = Depends(perms.require_page("pm")), db: Session = Depends(get_db), ym: str = ""):
    today = today_local()
    try:
        y, m = (int(x) for x in ym.split("-")) if ym else (today.year, today.month)
        date(y, m, 1)
    except Exception:
        y, m = today.year, today.month
    cal = pm_service.month_calendar(db, request.state.company.id, y, m)
    prev = f"{y - 1}-12" if m == 1 else f"{y}-{m - 1:02d}"
    nxt = f"{y + 1}-01" if m == 12 else f"{y}-{m + 1:02d}"
    return render(request, "pm_calendar.html", cal=cal, y=y, m=m, prev=prev, nxt=nxt, db=db)


@router.get("/pm/plan")
def pm_plan(request: Request, user: User = Depends(perms.require_page("pm")), db: Session = Depends(get_db), year: int = 0,
            line: str = "", customer: str = "", fmt: str = ""):
    company = request.state.company
    year = year or today_local().year
    plan = pm_service.year_plan(db, company.id, year, line, customer)
    if fmt == "xlsx":
        return _plan_excel(plan, company, line, customer)
    lines, custs = _lines_customers(db, company.id)
    return render(request, "pm_plan.html", plan=plan, year=year, line=line, customer=customer, lines=lines, customers=custs, db=db)


def _plan_excel(plan, company, line, customer):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    wb = Workbook()
    ws = wb.active
    ws.title = f"PM Plan {plan['year']}"
    ws.append([f"แผนการบำรุงรักษาแม่พิมพ์ประจำปี {plan['year']} (Die PM Annual Plan) — {company.name}"])
    ws.append([f"Line: {line or 'ทั้งหมด'} · ลูกค้า: {customer or 'ทั้งหมด'} · P = แผน, A = ทำจริง (วันที่ DD-MM)"])
    ws.append([])
    months = ["ม.ค.", "ก.พ.", "มี.ค.", "เม.ย.", "พ.ค.", "มิ.ย.", "ก.ค.", "ส.ค.", "ก.ย.", "ต.ค.", "พ.ย.", "ธ.ค."]
    head = ["Line", "ลูกค้า", "Part No", "Part Name", "Die Process", "Die Code", "รอบ (วัน)", "รอบ (stroke)", "", ] + months + ["แผน", "ทำจริง"]
    ws.append(head)
    hdr = ws.max_row
    thin = Side(style="thin", color="C9D3E3")
    for c in ws[hdr]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F5FBF")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    green, red, blue = PatternFill("solid", fgColor="D9F2E3"), PatternFill("solid", fgColor="FBE0E0"), PatternFill("solid", fgColor="E3ECFB")
    for r in plan["rows"]:
        d = r["die"]
        for kind in ("P", "A"):
            cells = []
            for m in r["months"]:
                if kind == "P":
                    cells.append(", ".join(x.strftime("%d") for x in m["p"]))
                else:
                    cells.append(", ".join(a["date"].strftime("%d") for a in m["a"]))
            base = [d.line, d.customer, d.part_no, d.part_name, d.die_process, d.die_code, d.pm_interval_days or "", d.pm_interval_strokes or ""] if kind == "P" else [""] * 8
            ws.append(base + [kind] + cells + ([r["planned"], r["actual"]] if kind == "P" else ["", ""]))
            row = ws.max_row
            for i, m in enumerate(r["months"]):
                c = ws.cell(row, 10 + i)
                c.alignment = Alignment(horizontal="center")
                c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
                if kind == "P" and m["p"]:
                    c.fill = red if m["late"] else blue
                if kind == "A" and m["a"]:
                    c.fill = green
    ws.append([])
    ws.append(["รวม", "", "", "", "", "", "", "", "", ] + [""] * 12 + [plan["planned"], plan["actual"]])
    for i, w in enumerate([10, 12, 18, 26, 22, 12, 9, 11, 4] + [7] * 12 + [7, 8], 1):
        ws.column_dimensions[ws.cell(hdr, i).column_letter].width = w
    ws.freeze_panes = ws.cell(hdr + 1, 10)
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename=PM_plan_{company.code}_{plan['year']}.xlsx"})


# ---------------------------------------------------------------- suggestions (PM interval from breakdown history)
@router.get("/pm/suggest")
def pm_suggest(request: Request, user: User = Depends(perms.require_page("pm")), db: Session = Depends(get_db)):
    rows = pm_service.suggestions(db, request.state.company.id, limit=100)
    return render(request, "pm_suggest.html", rows=rows, db=db)


@router.post("/pm/suggest/apply")
async def pm_suggest_apply(request: Request, user: User = Depends(require_roles(*LEADS)), db: Session = Depends(get_db)):
    form = await request.form()
    company = request.state.company
    n = 0
    for key in form.getlist("die_id"):
        die = db.get(Die, int(key))
        days = form.get(f"days_{key}") or ""
        if not die or die.company_id != company.id or not days.isdigit() or int(days) <= 0:
            continue
        old = die.pm_interval_days
        die.pm_interval_days = int(days)
        base = die.pm_last_date or today_local()
        die.pm_next_date = max(base + timedelta(days=die.pm_interval_days), today_local())
        db.add(PMEvent(company_id=company.id, die_id=die.id, kind="PLAN", from_date=None, to_date=die.pm_next_date,
                       reason=f"ตั้งรอบ PM {old or 0} → {die.pm_interval_days} วัน (จากประวัติเสีย)", user_name=user.full_name))
        n += 1
    audit(db, request, "pm_plan_apply", "die", n, "suggested intervals")
    db.commit()
    return RedirectResponse(f"/pm?msg=ตั้งแผน PM ให้ {n} แม่พิมพ์แล้ว", 303)


# ---------------------------------------------------------------- strokes
@router.get("/pm/strokes")
def strokes_page(request: Request, user: User = Depends(perms.require_page("pm_strokes")), db: Session = Depends(get_db), q: str = "",
                 die_id: int = 0):
    company = request.state.company
    die = db.get(Die, die_id) if die_id else None
    if die and die.company_id != company.id:
        die = None
    found = []
    if q and not die:
        like = f"%{q.strip()}%"
        found = (db.query(Die).filter(Die.company_id == company.id, or_(Die.part_no.ilike(like), Die.die_code.ilike(like),
                                                                        Die.part_name.ilike(like))).order_by(Die.part_no).limit(30).all())
        if len(found) == 1:
            die, found = found[0], []
    logs = (db.query(StrokeLog).filter(StrokeLog.company_id == company.id).order_by(StrokeLog.id.desc()).limit(40).all())
    stroke_dies = (db.query(Die).filter(Die.company_id == company.id, Die.is_active == True, Die.pm_interval_strokes > 0)
                   .order_by(Die.part_no).all())
    stroke_dies.sort(key=lambda d: -(d.stroke_pct or 0))
    return render(request, "pm_strokes.html", die=die, found=found, q=q, logs=logs, stroke_dies=stroke_dies[:60], db=db)


@router.post("/pm/strokes")
def strokes_save(request: Request, user: User = Depends(perms.require_page("pm_strokes")), db: Session = Depends(get_db),
                 die_id: int = Form(...), value: str = Form(...), mode: str = Form("add"), work_date: str = Form(""), note: str = Form("")):
    company = request.state.company
    die = db.get(Die, die_id)
    if not die or die.company_id != company.id:
        raise HTTPException(404)
    v = value.replace(",", "").strip()
    if not v.isdigit():
        return RedirectResponse(f"/pm/strokes?die_id={die.id}&msg=กรุณาใส่ตัวเลข", 303)
    try:
        pm_service.record_strokes(db, company.id, die, int(v), mode, "MANUAL", user.full_name, _date(work_date), note)
    except ValueError as e:
        return RedirectResponse(f"/pm/strokes?die_id={die.id}&msg={e}", 303)
    db.commit()
    return RedirectResponse(f"/pm/strokes?die_id={die.id}&msg=บันทึกแล้ว ยอดสะสม {die.stroke_count:,} stroke", 303)


STROKE_COLS = ["DieCode", "PartNo", "DieProcess", "Strokes", "Mode (add/total)", "Date (DD-MM-YYYY)", "Note"]


@router.get("/pm/strokes/template.xlsx")
def strokes_template(request: Request, user: User = Depends(perms.require_page("pm_strokes"))):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Strokes"
    ws.append(STROKE_COLS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F5FBF")
    ws.append(["D-0001", "", "", 12500, "add", today_local().strftime("%d-%m-%Y"), "กะเช้า"])
    ws.append(["", "44591-0K110", "OP10 Progressive", 350000, "total", today_local().strftime("%d-%m-%Y"), "อ่านจาก counter"])
    for col, w in zip("ABCDEFG", (14, 18, 24, 12, 16, 18, 24)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": "attachment; filename=stroke_upload_template.xlsx"})


@router.post("/pm/strokes/upload")
async def strokes_upload(request: Request, user: User = Depends(perms.require_page("pm_strokes")), db: Session = Depends(get_db),
                         file: UploadFile = File(...)):
    from ..services.import_data import read_rows
    company = request.state.company
    rows = read_rows(file.filename, await file.read())
    ok, errs = 0, []
    for i, raw in enumerate(rows, 2):
        r = {str(k).split(" ")[0].strip().lower(): ("" if v is None else str(v).strip()) for k, v in raw.items() if k}
        die = pm_service.find_die(db, company.id, r.get("diecode"), r.get("partno"), r.get("dieprocess"))
        if not die:
            errs.append(f"แถว {i}: ไม่พบแม่พิมพ์")
            continue
        val = r.get("strokes", "").replace(",", "").split(".")[0]
        if not val.isdigit():
            errs.append(f"แถว {i}: Strokes ต้องเป็นตัวเลข")
            continue
        try:
            wd = _date(r.get("date", "")[:10]) if r.get("date") else None
        except HTTPException:
            wd = None
        try:
            pm_service.record_strokes(db, company.id, die, int(val), "total" if r.get("mode", "").lower().startswith("t") else "add",
                                      "EXCEL", user.full_name, wd, r.get("note", ""))
            ok += 1
        except ValueError as e:
            errs.append(f"แถว {i}: {e}")
    audit(db, request, "strokes_upload", "die", file.filename, f"{ok} ok, {len(errs)} errors")
    db.commit()
    msg = f"บันทึก stroke {ok} แถว" + (f" · ปัญหา {len(errs)}: {'; '.join(errs[:5])}" if errs else "")
    return RedirectResponse(f"/pm/strokes?msg={msg}", 303)


# ---------------------------------------------------------------- checklist templates
@router.get("/pm/templates")
def templates_page(request: Request, user: User = Depends(perms.require_page("pm")), db: Session = Depends(get_db), tid: int = 0):
    company = request.state.company
    pm_service.ensure_templates(db, company.id)
    db.commit()
    temps = db.query(PMTemplate).filter(PMTemplate.company_id == company.id).order_by(PMTemplate.id).all()
    cur = next((t for t in temps if t.id == tid), temps[0] if temps else None)
    usage = {t.id: db.query(Die).filter(Die.company_id == company.id, Die.pm_template_id == t.id).count() for t in temps}
    can_edit = user.is_superadmin or user.role in LEADS
    return render(request, "pm_templates.html", temps=temps, cur=cur, usage=usage, can_edit=can_edit, db=db)


def _tmpl(db, request, tid):
    t = db.get(PMTemplate, tid)
    if not t or t.company_id != request.state.company.id:
        raise HTTPException(404)
    return t


def _f(v):
    try:
        return float(v) if str(v).strip() != "" else None
    except ValueError:
        return None


@router.post("/pm/templates")
def template_save(request: Request, user: User = Depends(require_roles(*LEADS)), db: Session = Depends(get_db), tid: int = Form(0),
                  name: str = Form(...), match: str = Form(""), is_active: str = Form("1"), copy_from: int = Form(0)):
    company = request.state.company
    t = _tmpl(db, request, tid) if tid else PMTemplate(company_id=company.id)
    t.name, t.match, t.is_active = name.strip()[:120], match.strip()[:120], is_active == "1"
    if not tid:
        db.add(t)
        if copy_from:
            src = _tmpl(db, request, copy_from)
            for it in src.items:
                t.items.append(PMTemplateItem(sort=it.sort, section=it.section, text=it.text, kind=it.kind, unit=it.unit,
                                              spec_min=it.spec_min, spec_max=it.spec_max))
    audit(db, request, "pm_template_save", "pm_template", t.id or 0, t.name)
    db.commit()
    return RedirectResponse(f"/pm/templates?tid={t.id}&msg=บันทึกแล้ว", 303)


@router.post("/pm/templates/{tid}/delete")
def template_delete(request: Request, tid: int, user: User = Depends(require_roles(*LEADS)), db: Session = Depends(get_db)):
    t = _tmpl(db, request, tid)
    db.query(Die).filter(Die.pm_template_id == t.id).update({Die.pm_template_id: None}, synchronize_session=False)
    audit(db, request, "pm_template_delete", "pm_template", t.id, t.name)
    db.delete(t)
    db.commit()
    return RedirectResponse("/pm/templates?msg=ลบแล้ว", 303)


@router.post("/pm/templates/{tid}/items")
def template_item_save(request: Request, tid: int, user: User = Depends(require_roles(*LEADS)), db: Session = Depends(get_db),
                       iid: int = Form(0), section: str = Form(""), text: str = Form(...), kind: str = Form("check"), unit: str = Form(""),
                       spec_min: str = Form(""), spec_max: str = Form(""), sort: str = Form("")):
    t = _tmpl(db, request, tid)
    it = next((x for x in t.items if x.id == iid), None) if iid else PMTemplateItem()
    if it is None:
        raise HTTPException(404)
    it.section, it.text = section.strip()[:60], text.strip()[:200]
    it.kind = "value" if kind == "value" else "check"
    it.unit, it.spec_min, it.spec_max = unit.strip()[:16], _f(spec_min), _f(spec_max)
    it.sort = int(sort) if sort.strip().lstrip("-").isdigit() else (max([x.sort for x in t.items] or [0]) + 1)
    if not iid:
        t.items.append(it)
    db.commit()
    return RedirectResponse(f"/pm/templates?tid={t.id}", 303)


@router.post("/pm/templates/{tid}/items/{iid}/delete")
def template_item_delete(request: Request, tid: int, iid: int, user: User = Depends(require_roles(*LEADS)), db: Session = Depends(get_db)):
    t = _tmpl(db, request, tid)
    it = next((x for x in t.items if x.id == iid), None)
    if it:
        t.items.remove(it)
        db.commit()
    return RedirectResponse(f"/pm/templates?tid={t.id}", 303)


@router.post("/dies/{die_id}/pm-plan")
def die_pm_plan(request: Request, die_id: int, user: User = Depends(require_roles(*LEADS, "foreman")), db: Session = Depends(get_db),
                pm_interval_days: str = Form("0"), pm_interval_strokes: str = Form("0"), pm_last_date: str = Form(""),
                pm_template_id: str = Form(""), stroke_count: str = Form("")):
    die = db.get(Die, die_id)
    if not die or die.company_id != request.state.company.id:
        raise HTTPException(404)
    num = lambda v: int(str(v).replace(",", "").strip() or 0) if str(v).replace(",", "").strip().isdigit() else 0
    old = (die.pm_interval_days, die.pm_interval_strokes)
    die.pm_interval_days, die.pm_interval_strokes = num(pm_interval_days), num(pm_interval_strokes)
    die.pm_last_date = _date(pm_last_date)
    die.pm_next_date = (die.pm_last_date + timedelta(days=die.pm_interval_days)) if die.pm_last_date and die.pm_interval_days else (
        today_local() if die.pm_interval_days and not die.pm_last_date else None)
    tid = int(pm_template_id) if pm_template_id.isdigit() else None
    if tid:
        _tmpl(db, request, tid)
    die.pm_template_id = tid
    if stroke_count.strip():
        v = num(stroke_count)
        if v != (die.stroke_count or 0):
            try:
                pm_service.record_strokes(db, die.company_id, die, v, "total", "MANUAL", user.full_name, note="แก้ที่หน้าแม่พิมพ์")
            except ValueError:
                die.stroke_count = v           # counter reset / correction by a leader
                die.strokes_at_last_pm = min(die.strokes_at_last_pm or 0, v)
    if old != (die.pm_interval_days, die.pm_interval_strokes):
        db.add(PMEvent(company_id=die.company_id, die_id=die.id, kind="PLAN", to_date=die.pm_next_date,
                       reason=f"รอบ PM {old[0]}→{die.pm_interval_days} วัน, {old[1] or 0}→{die.pm_interval_strokes} stroke", user_name=user.full_name))
    audit(db, request, "pm_plan", "die", die.id, f"{die.pm_interval_days}d {die.pm_interval_strokes}st")
    db.commit()
    return RedirectResponse(f"/dies/{die.id}?msg=บันทึกแผน PM แล้ว", 303)


# ---------------------------------------------------------------- checklist on a PM job
CHECK_EDIT_STATUSES = (ST["ASSIGNED"], ST["IN_PROGRESS"], ST["RETURNED"], ST["NG_RETURN"])


def can_fill_checklist(user, req) -> bool:
    if req.request_type != "PM" or req.status not in CHECK_EDIT_STATUSES:
        return False
    if user.is_superadmin or user.role in ("admin", "die_leader"):
        return True
    return user.role == "technician" and any(a.user_id == user.id for a in req.assignments)


@router.post("/requests/{rid}/checklist")
async def checklist_save(request: Request, rid: int, user: User = Depends(require_user), db: Session = Depends(get_db)):
    from .requests import get_req
    req = get_req(db, request, rid)
    if not can_fill_checklist(user, req):
        raise HTTPException(403, "บันทึก checklist ได้เฉพาะช่างที่ได้รับมอบหมาย / หัวหน้า Die ระหว่างทำงาน")
    form = await request.form()
    now = now_local()
    for c in req.pm_checks:
        res = (form.get(f"r_{c.id}") or "").upper()
        val = _f(form.get(f"v_{c.id}") or "")
        note = (form.get(f"n_{c.id}") or "").strip()[:200]
        if c.kind == "value":
            c.value = val
            res = pm_service.judge_value(c) or ("NA" if res == "NA" else ("" if val is None else "OK"))
        if res in ("OK", "NG", "NA") and (res != c.result or note != (c.note or "") or c.kind == "value"):
            c.result, c.note, c.user_name, c.at = res, note or None, user.full_name, now
        elif not res:
            c.note = note or c.note
    stt = pm_service.checklist_state(req)
    if stt["answered"]:
        req.progress_pct = max(req.progress_pct or 0, min(90, 10 + int(stt["pct"] * 0.8)))
    audit(db, request, "pm_checklist", "request", req.id, f"{stt['answered']}/{stt['total']} NG {stt['ng']}")
    db.commit()
    return RedirectResponse(f"/requests/{req.id}?msg=บันทึก checklist {stt['answered']}/{stt['total']} รายการ#pmcheck", 303)


@router.post("/requests/{rid}/checklist/reset")
def checklist_reset(request: Request, rid: int, user: User = Depends(require_user), db: Session = Depends(get_db)):
    from .requests import get_req
    req = get_req(db, request, rid)
    if not can_fill_checklist(user, req) or user.role == "technician":
        raise HTTPException(403)
    n = pm_service.attach_checklist(db, req, force=True)
    db.commit()
    return RedirectResponse(f"/requests/{req.id}?msg=โหลด checklist ใหม่ {n} รายการ#pmcheck", 303)


@router.get("/requests/{rid}/pm-print")
def pm_print(request: Request, rid: int, user: User = Depends(require_user), db: Session = Depends(get_db)):
    from .requests import get_req
    req = get_req(db, request, rid)
    if req.request_type != "PM":
        raise HTTPException(404)
    return render(request, "pm_print.html", req=req, st8=pm_service.checklist_state(req), db=db)


# ---------------------------------------------------------------- die register
@router.get("/dies")
def dies(request: Request, user: User = Depends(require_user), db: Session = Depends(get_db), q: str = "", customer: str = "",
         line: str = "", pm: str = "", page: int = 1):
    company = request.state.company
    qry = db.query(Die).filter(Die.company_id == company.id)
    if q:
        like = f"%{q}%"
        qry = qry.filter(or_(Die.part_no.ilike(like), Die.part_name.ilike(like), Die.die_code.ilike(like), Die.die_process.ilike(like), Die.model.ilike(like)))
    if customer:
        qry = qry.filter(Die.customer == customer)
    if line:
        qry = qry.filter(Die.line == line)
    if pm == "yes":
        qry = qry.filter(Die.pm_interval_days > 0)
    total = qry.count()
    per = 50
    rows = qry.order_by(Die.customer, Die.part_no, Die.die_process).offset((page - 1) * per).limit(per).all()
    customers = sorted({c[0] for c in db.query(Die.customer).filter(Die.company_id == company.id).distinct() if c[0]})
    lines = sorted({l[0] for l in db.query(Die.line).filter(Die.company_id == company.id).distinct() if l[0]})
    return render(request, "dies.html", rows=rows, total=total, page=page, pages=max(1, -(-total // per)), q=q, customer=customer,
                  line=line, pm=pm, customers=customers, lines=lines, db=db)


@router.get("/dies/{die_id}")
def die_detail(request: Request, die_id: int, user: User = Depends(require_user), db: Session = Depends(get_db)):
    company = request.state.company
    die = db.get(Die, die_id)
    if not die or die.company_id != company.id:
        raise HTTPException(404)
    hist = perms.scope_requests(db.query(RepairRequest), user).filter(RepairRequest.die_id == die.id, RepairRequest.is_deleted == False).order_by(RepairRequest.created_at.desc()).all()
    gaps = []
    ts = sorted(r.created_at for r in hist if r.request_type == "REPAIR")
    gaps = [(b - a).days for a, b in zip(ts, ts[1:])]
    mtbf = round(sum(gaps) / len(gaps), 1) if gaps else None
    net = [r.net_repair_minutes for r in hist if r.net_repair_minutes]
    mttr = round(sum(net) / len(net)) if net else None
    cost = round(sum(r.parts_cost or 0 for r in hist), 2)
    problems = {}
    for r in hist:
        problems[r.problem_text or "-"] = problems.get(r.problem_text or "-", 0) + 1
    pm_service.ensure_templates(db, company.id)
    pct = int(company.settings.get("pm_due_stroke_pct", 90) or 90)
    return render(request, "die_detail.html", die=die, hist=hist, mtbf=mtbf, mttr=mttr, cost=cost,
                  problems=sorted(problems.items(), key=lambda x: -x[1]),
                  pm_st=pm_service.status(die, today_local(), 14, pct), tmpl=pm_service.template_for(db, die),
                  temps=db.query(PMTemplate).filter(PMTemplate.company_id == company.id, PMTemplate.is_active == True).all(),
                  events=db.query(PMEvent).filter(PMEvent.die_id == die.id).order_by(PMEvent.id.desc()).limit(15).all(),
                  strokes=db.query(StrokeLog).filter(StrokeLog.die_id == die.id).order_by(StrokeLog.id.desc()).limit(10).all(),
                  has_open_pm=die.id in pm_service.open_pm_die_ids(db, company.id), db=db)


@router.post("/dies/{die_id}")
def die_save(request: Request, die_id: int, user: User = Depends(require_roles("die_leader", "admin", "manager", "foreman")),
             db: Session = Depends(get_db), die_code: str = Form(""), customer: str = Form(""), model: str = Form(""),
             part_name: str = Form(""), die_process: str = Form(""), line: str = Form(""), machine_no: str = Form(""),
             pm_interval_days: str = Form(None), pm_last_date: str = Form(None), is_active: str = Form("1")):
    die = db.get(Die, die_id)
    if not die or die.company_id != request.state.company.id:
        raise HTTPException(404)
    die.die_code, die.customer, die.model, die.part_name = die_code.strip(), customer.strip(), model.strip(), part_name.strip()
    die.die_process, die.line, die.machine_no = die_process.strip() or die.die_process, line.strip(), machine_no.strip()
    if pm_interval_days is not None:                 # older forms / API still send the plan here
        die.pm_interval_days = max(0, int(pm_interval_days) if str(pm_interval_days).isdigit() else 0)
    if pm_last_date is not None:
        die.pm_last_date = _date(pm_last_date)
        die.pm_next_date = (die.pm_last_date + timedelta(days=die.pm_interval_days)) if die.pm_last_date and die.pm_interval_days else None
    die.is_active = is_active == "1"
    audit(db, request, "edit", "die", die.id, die.label)
    db.commit()
    return RedirectResponse(f"/dies/{die.id}", 303)


@router.post("/dies")
def die_create(request: Request, user: User = Depends(require_roles("die_leader", "admin", "manager", "foreman")), db: Session = Depends(get_db),
               die_code: str = Form(""), customer: str = Form(""), model: str = Form(""), part_no: str = Form(...), part_name: str = Form(""),
               die_process: str = Form(...), line: str = Form(""), pm_interval_days: int = Form(0)):
    die = Die(company_id=request.state.company.id, die_code=die_code.strip(), customer=customer.strip(), model=model.strip(),
              part_no=part_no.strip(), part_name=part_name.strip(), die_process=die_process.strip(), line=line.strip(),
              pm_interval_days=max(0, pm_interval_days))
    db.add(die)
    db.commit()
    return RedirectResponse(f"/dies/{die.id}", 303)
