import csv
import io
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..config import BASE_DIR
from ..i18n import t
from ..database import now_local

COLUMNS = [
    ("dr_no", "DR No."), ("legacy_id", "Legacy ID"), ("request_type", "Type"), ("status", "Status"),
    ("created_at", "Request Date"), ("requester_name", "Requester"), ("line", "Line"), ("customer", "Customer"),
    ("model", "Model"), ("part_no", "Part No."), ("part_name", "Part Name"), ("die_process", "Die Process"),
    ("die_code", "Die Code"), ("required_day", "Required Day"), ("is_emergency", "Emergency"),
    ("target_date", "Target Date"), ("stock_oh", "Stock OH"), ("problem_text", "Problem"), ("other_details", "Other Details"),
    ("foreman_name", "Foreman"), ("foreman_at", "Foreman Approve"), ("leader_name", "Die Leader"), ("leader_at", "Leader Approve"),
    ("technician_names", "Technicians"), ("tech_received_at", "Received"), ("tech_completed_at", "Completed"),
    ("net_repair_minutes", "Net Repair (min)"), ("root_cause_text", "Root Cause"), ("action_text", "Action"),
    ("tryout_count", "Try-out #"), ("tryout_result", "Try-out Result"), ("tryout_at", "Try-out Date"),
    ("qc_result", "QC Result"), ("qc_comment", "QC Comment"), ("qc_at", "QC Date"), ("ng_count", "NG Count"),
    ("closed_at", "Closed"), ("downtime_minutes", "Downtime (min)"), ("parts_cost", "Parts Cost"),
]


_YM = __import__("re").compile(r"^(\d{4})-(\d{2})$")
_YMD = __import__("re").compile(r"^(\d{4})-(\d{2})-(\d{2})$")


def dmy_text(v):
    """Text cells like '2026-10' / '2026-10-16' are shown day/month first: '10-2026' / '16-10-2026'."""
    if isinstance(v, str):
        m = _YMD.match(v)
        if m:
            return f"{m.group(3)}-{m.group(2)}-{m.group(1)}"
        m = _YM.match(v)
        if m:
            return f"{m.group(2)}-{m.group(1)}"
    return v


def _val(r, key):
    if key == "foreman_name":
        return r.foreman.full_name if r.foreman else ""
    if key == "leader_name":
        return r.leader.full_name if r.leader else ""
    v = getattr(r, key, "")
    if isinstance(v, datetime):
        return v.strftime("%d-%m-%Y %H:%M")
    if isinstance(v, date):
        return v.strftime("%d-%m-%Y")
    if isinstance(v, bool):
        return "Y" if v else ""
    return v if v is not None else ""


def requests_to_rows(reqs):
    return [[_val(r, k) for k, _ in COLUMNS] for r in reqs]


def to_csv(reqs) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([h for _, h in COLUMNS])
    w.writerows(requests_to_rows(reqs))
    return ("﻿" + buf.getvalue()).encode("utf-8")


def to_xlsx(reqs, title="Die Requests", extra_sheets: dict = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Requests"
    ws.append([h for _, h in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F5FBF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in requests_to_rows(reqs):
        ws.append(row)
    for i, (_, h) in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, min(40, len(h) + 6))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for name, (headers, rows) in (extra_sheets or {}).items():
        s = wb.create_sheet(name[:30])
        s.append([dmy_text(h) for h in headers])
        for c in s[1]:
            c.font = Font(bold=True)
        for r in rows:
            s.append([dmy_text(v) for v in r])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _esc(v) -> str:
    """ReportLab Paragraph parses mini-HTML: '<', '>' and '&' typed by users would crash the PDF."""
    import html
    return html.escape(str(v), quote=False)


def _thai_font():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    if "ThaiFont" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("ThaiFont", str(BASE_DIR / "static/fonts/FreeSerif.ttf")))
        pdfmetrics.registerFont(TTFont("ThaiFontBold", str(BASE_DIR / "static/fonts/FreeSerifBold.ttf")))
    return "ThaiFont"


def to_pdf_table(title: str, headers, rows, kpis: dict = None, landscape_mode=True) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    font = _thai_font()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4) if landscape_mode else A4, leftMargin=20, rightMargin=20,
                            topMargin=24, bottomMargin=20)
    h1 = ParagraphStyle("h1", fontName="ThaiFontBold", fontSize=15, leading=19)
    body = ParagraphStyle("b", fontName=font, fontSize=8.5, leading=11)
    story = [Paragraph(_esc(title), h1), Paragraph(f"พิมพ์เมื่อ {now_local():%d-%m-%Y %H:%M}", body), Spacer(1, 8)]
    if kpis:
        story.append(Paragraph("  |  ".join(f"<b>{_esc(k)}</b>: {_esc(v)}" for k, v in kpis.items()), body))
        story.append(Spacer(1, 8))
    data = [[Paragraph(_esc(h), body) for h in headers]] + [[Paragraph(_esc(c), body) for c in r] for r in rows]
    tbl = Table(data, repeatRows=1)
    tbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F5FBF")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F6FB")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.append(tbl)
    doc.build(story)
    return buf.getvalue()


def to_pdf_requests(reqs, title="Die Repair / PM Report", kpis=None) -> bytes:
    keys = ["dr_no", "request_type", "status", "created_at", "line", "customer", "part_no", "die_process",
            "problem_text", "technician_names", "net_repair_minutes", "qc_result", "closed_at"]
    ordered = [(k, h) for k, h in COLUMNS if k in keys]
    heads = [h for _, h in ordered]
    rows = [[_val(r, k) for k, _ in ordered] for r in reqs]
    return to_pdf_table(title, heads, rows, kpis)
