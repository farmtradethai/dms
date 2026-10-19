"""New-company start-up kit.

ONE Excel workbook holds every piece of starting data a company needs before go-live:
company settings, lines, customers, problems, root causes, users, approval routing, die register,
and (optional) old job history.  Upload -> every sheet is validated first (nothing saved), the admin sees
row-by-row errors, then imports for real.  `readiness()` lists what is still missing before go-live."""
import io
import re
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from ..models import (ROLES, ApproverRule, Company, Customer, Die, NotificationChannel, Problem, ProdLine, RootCause, User)
from .import_data import TEMPLATE_COLS, clean, norm_line

YES = ("Y", "YES", "1", "TRUE", "ใช่", "ใช้งาน")
STEP_CODES = ["FOREMAN", "LEADER", "TRYOUT", "QC"]
SHIFT_KEYS = ("DAY", "NIGHT")

COMPANY_KEYS = [
    # key, label, example, help
    ("company_name", "ชื่อบริษัท", "AAPICO Hitech Public Co., Ltd.", "ชื่อเต็มที่แสดงในระบบ/ใบพิมพ์"),
    ("logo_text", "ข้อความโลโก้", "AAPICO", ""),
    ("dr_prefix", "Prefix เลขใบงาน", "AHMD", "เช่น AHMD → AHMD00001 หรือ AHR-DR"),
    ("dr_style", "รูปแบบเลขใบงาน", "running", "running = ต่อเนื่อง (AHMD05001) | monthly = รายเดือน (DR-2610-001)"),
    ("dr_start_no", "เลขใบงานล่าสุดที่ใช้ไปแล้ว", 0, "ใบแรกในระบบใหม่จะเป็นเลขนี้ +1 (ต่อจากระบบเดิม)"),
    ("emergency_required_day", "Required Day ที่นับเป็นงานด่วน", "<1", "<1 / 3 / 5 / >7"),
    ("approval_reminder_hours", "เตือนผู้อนุมัติเมื่อค้างเกิน (ชม.)", 8, "0 = ไม่เตือน"),
    ("pm_skip_foreman", "งาน PM ข้ามขั้น Foreman", "N", "Y / N"),
    ("approver_visibility", "ผู้อนุมัติเห็นใบงาน", "own", "own = เฉพาะ Line ที่ตนอนุมัติ | all = ทั้งบริษัท"),
    ("DAY_start", "กะเช้า เริ่ม", "08:00", "HH:MM"),
    ("DAY_end", "กะเช้า สิ้นสุด", "20:00", "HH:MM"),
    ("DAY_breaks", "กะเช้า เวลาพัก (ไม่นับเวลาซ่อม)", "10:00-10:10, 12:00-13:00, 15:00-15:10, 17:00-17:30", "ช่วงเวลาคั่นด้วย ,"),
    ("NIGHT_start", "กะดึก เริ่ม", "20:00", "HH:MM"),
    ("NIGHT_end", "กะดึก สิ้นสุด", "08:00", "HH:MM"),
    ("NIGHT_breaks", "กะดึก เวลาพัก", "22:00-22:10, 00:00-01:00, 03:00-03:10, 05:00-05:30", "ช่วงเวลาคั่นด้วย ,"),
    ("tv_working_hours", "เวลาที่จอ TV อัปเดต", "", "เช่น 07:30-20:30 (ว่าง = ตลอด)"),
    ("spare_app_url", "URL โปรแกรม Die Spare Part", "", "ใช้ลิงก์ QR ไปโปรแกรมอะไหล่"),
]

# sheet name -> (title, [(key, header, required, width, help)], example rows)
SHEETS = {
    "2_Lines": ("Line การผลิต", [("line", "Line *", True, 22, "ชื่อ Line เช่น Line3, Line4-5"), ("active", "ใช้งาน (Y/N)", False, 14, "")],
                [["Line3", "Y"], ["Line4-5", "Y"]]),
    "3_Customers": ("ลูกค้า", [("customer", "Customer *", True, 26, ""), ("active", "ใช้งาน (Y/N)", False, 14, "")],
                    [["TOYOTA", "Y"], ["ISUZU", "Y"]]),
    "4_Problems": ("ปัญหา (Production เลือกตอนแจ้ง)", [("name_th", "ปัญหา (ไทย) *", True, 30, ""), ("name_en", "Problem (EN)", False, 30, ""),
                                                       ("sort", "ลำดับ", False, 10, "")],
                   [["ครีบรู", "Hole burr", 1], ["อื่นๆ", "Others", 99]]),
    "5_RootCauses": ("สาเหตุ (ช่างเลือกตอนปิดงาน)", [("name_th", "สาเหตุ (ไทย) *", True, 34, ""), ("name_en", "Root cause (EN)", False, 30, ""),
                                                     ("sort", "ลำดับ", False, 10, "")],
                     [["Punchสึก+Buttonสึก", "Punch & button worn", 1], ["อื่นๆ", "Others", 99]]),
    "6_Users": ("ผู้ใช้", [("username", "username *", True, 18, "ภาษาอังกฤษ ไม่มีช่องว่าง"), ("full_name", "full_name *", True, 26, "ต้องตรงกับชื่อในข้อมูลเก่า (ถ้ามี)"),
                          ("role", "role *", True, 14, ""), ("password", "password", False, 14, "ผู้ใช้ใหม่ต้องใส่ · ว่าง = ไม่เปลี่ยน"),
                          ("emp_code", "emp_code", False, 12, ""), ("position", "position", False, 16, ""), ("department", "department", False, 16, ""),
                          ("line", "line", False, 12, "Line ประจำ (ว่าง = ทุก Line)"), ("shift", "shift", False, 10, "DAY/NIGHT/A/B"),
                          ("email", "email", False, 26, "ผู้อนุมัติควรมี (รับอีเมลแจ้งอนุมัติ)"), ("phone", "phone", False, 14, ""),
                          ("skills", "skills", False, 20, "ช่าง: เชื่อม, เจียร"), ("is_active", "is_active (Y/N)", False, 12, "")],
                [["admin.ah", "ผู้ดูแลระบบ AH", "admin", "Temp#2026", "", "Admin", "IT", "", "", "admin@aapico.com", "", "", "Y"],
                 ["prod.l3", "นาย ผลิต สายสาม", "production", "Temp#2026", "10000001", "Line Leader", "Production", "Line3", "DAY", "", "", "", "Y"],
                 ["fm.somsak", "นาย สมศักดิ์ อนุมัติ", "foreman", "Temp#2026", "10000002", "Foreman", "Production", "Line3", "DAY", "somsak@aapico.com", "", "", "Y"],
                 ["dl.panat", "นาย พนัด หัวหน้า", "die_leader", "Temp#2026", "10000003", "Die Leader", "Die Maintenance", "", "DAY", "panat@aapico.com", "", "", "Y"],
                 ["tech.a", "นาย ช่าง หนึ่ง", "technician", "Temp#2026", "10000004", "Technician", "Die Maintenance", "", "DAY", "", "", "เชื่อม, เจียร", "Y"],
                 ["qc.l3", "นางสาว คิวซี ตรวจ", "qc", "Temp#2026", "10000005", "QC", "QC", "Line3", "DAY", "qc@aapico.com", "", "", "Y"]]),
    "7_Approvers": ("ผังผู้อนุมัติ", [("step", "step *", True, 12, "FOREMAN / LEADER / TRYOUT / QC"), ("line", "line", False, 14, "ว่าง = ทุก Line"),
                                     ("username", "username *", True, 18, "ต้องมีในชีต 6_Users หรือในระบบแล้ว"), ("backup", "สำรอง (Y/N)", False, 12, "Y = ผู้อนุมัติแทน")],
                    [["FOREMAN", "Line3", "fm.somsak", "N"], ["LEADER", "", "dl.panat", "N"], ["QC", "Line3", "qc.l3", "N"], ["TRYOUT", "Line3", "prod.l3", "N"]]),
    "8_Dies": ("ทะเบียนแม่พิมพ์", [("customer", "Customer", False, 14, ""), ("model", "Model", False, 12, ""), ("part_no", "PartNo *", True, 18, ""),
                                   ("part_name", "PartName", False, 30, ""), ("die_process", "DieProcess *", True, 22, "เช่น OP10 Progressive"),
                                   ("die_code", "DieCode", False, 14, "รหัส/QR บนแม่พิมพ์"), ("line", "Line", False, 12, ""), ("machine_no", "Machine", False, 12, ""),
                                   ("pm_interval_days", "PM Interval Days", False, 14, "0 = ไม่มีแผน PM"),
                                   ("pm_interval_strokes", "PM Interval Strokes", False, 16, "PM ทุกกี่ stroke (0 = ไม่ใช้)"),
                                   ("stroke_count", "Stroke Count", False, 14, "Stroke สะสมปัจจุบัน"),
                                   ("pm_last_date", "PM Last Date (DD-MM-YYYY)", False, 18, "PM ครั้งล่าสุด"), ("active", "Active (Y/N)", False, 10, "")],
               [["TOYOTA", "640A", "44591-0K110", "BRACKET BRAKE ACTUATOR NO.1", "OP10 Progressive", "D-0001", "Line3", "P-800T", 90, 150000, 0, "01-09-2026", "Y"]]),
}
HISTORY_SHEET = "9_History"


# ================================================================== template
def _style_header(ws, cols, req_flags):
    for i, (cell, req) in enumerate(zip(ws[1], req_flags)):
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="C0392B" if req else "1F5FBF")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"


def build_kit(db=None, company=None, with_data=False) -> bytes:
    wb = Workbook()
    info = wb.active
    info.title = "0_อ่านก่อน"
    guide = [
        ("Die Management System — ชุดข้อมูลตั้งต้นสำหรับบริษัทใหม่", True),
        ("", False),
        ("วิธีใช้", True),
        ("1) กรอกทุกชีตตามลำดับ (หัวคอลัมน์สีแดง = ต้องกรอก, สีน้ำเงิน = ไม่บังคับ) · แถวตัวอย่างลบ/แก้ทับได้", False),
        ("2) อัปโหลดทั้งไฟล์ที่ Admin › 🚀 เริ่มต้นบริษัท › 'ตรวจไฟล์' — ระบบตรวจทุกแถวก่อน ยังไม่บันทึกอะไร", False),
        ("3) แก้ตามรายการข้อผิดพลาด แล้วอัปโหลดซ้ำ จนผ่าน → กด 'นำเข้าจริง'", False),
        ("4) ดู 'ความพร้อมก่อนใช้งานจริง' ในหน้าเดียวกัน ต้องผ่านทุกข้อ 'จำเป็น' ก่อนเปิดใช้", False),
        ("5) อัปโหลดซ้ำได้เสมอ: แถวที่มีอยู่แล้วจะถูกอัปเดต (ไม่สร้างซ้ำ) — ใช้เป็นวิธีแก้ข้อมูลทีละมากได้", False),
        ("", False),
        ("ชีต", True),
        ("1_Company  — ตั้งค่าบริษัท: เลขใบงาน, เวลากะ/เวลาพัก (หักออกจากเวลาซ่อม), เกณฑ์งานด่วน", False),
        ("2_Lines / 3_Customers — รายการ Line และลูกค้า (ใช้ใน dropdown)", False),
        ("4_Problems / 5_RootCauses — ปัญหาที่ Production เลือก / สาเหตุที่ช่างเลือก", False),
        ("6_Users — ทุกคนที่ใช้ระบบ (role: " + ", ".join(r for r in ROLES if r != "superadmin") + ")", False),
        ("7_Approvers — ผังผู้อนุมัติ: แต่ละขั้น (FOREMAN / LEADER / TRYOUT / QC) ของแต่ละ Line ใครเป็นผู้อนุมัติ, line ว่าง = ทุก Line", False),
        ("8_Dies — ทะเบียนแม่พิมพ์ (Part No + Process = 1 แม่พิมพ์) + รอบ PM", False),
        ("9_History — (ไม่บังคับ) ใบแจ้งซ่อมเก่าย้อนหลัง หรือใช้ไฟล์ DieRepairList.csv จาก Power Apps ที่หน้า นำเข้าข้อมูล", False),
        ("", False),
        ("ข้อควรรู้", True),
        ("• วันที่ใช้รูปแบบ DD-MM-YYYY · ใช้งาน/สำรอง ใส่ Y หรือ N", False),
        ("• ผู้ใช้ใหม่จะถูกบังคับเปลี่ยนรหัสผ่านเมื่อเข้าครั้งแรก", False),
        ("• ถ้าขั้นไหนไม่มีผู้อนุมัติในชีต 7 ระบบใช้ 'บทบาท + Line ประจำ' ของผู้ใช้แทน (ดูผลได้ที่ Admin › ผังผู้อนุมัติ)", False),
        ("• ข้อมูลของแต่ละบริษัทแยกกันสิ้นเชิง ไฟล์นี้จะนำเข้าเฉพาะบริษัทที่คุณกำลังใช้งานอยู่", False),
    ]
    for text, bold in guide:
        info.append([text])
        if bold:
            info.cell(info.max_row, 1).font = Font(bold=True, size=13 if info.max_row == 1 else 11, color="1F5FBF")
    info.column_dimensions["A"].width = 130

    # 1_Company
    ws = wb.create_sheet("1_Company")
    ws.append(["Key", "รายการ", "Value", "คำอธิบาย"])
    _style_header(ws, 4, [True, False, True, False])
    s = company.settings if company is not None else {}
    for key, label, example, hlp in COMPANY_KEYS:
        val = example
        if with_data and company is not None:
            val = _company_value(company, key, example)
        ws.append([key, label, val, hlp])
    for col, w in zip("ABCD", (24, 34, 50, 60)):
        ws.column_dimensions[col].width = w

    for name, (title, cols, examples) in SHEETS.items():
        ws = wb.create_sheet(name)
        ws.append([c[1] for c in cols])
        _style_header(ws, len(cols), [c[2] for c in cols])
        for i, c in enumerate(cols, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = c[3]
            if c[4]:
                ws.cell(1, i).comment = Comment(c[4], "DMS")
        rows = _current_rows(db, company, name) if (with_data and company is not None) else None
        for r in (rows if rows else examples):
            ws.append(r)
        _validations(ws, name, cols)

    ws = wb.create_sheet(HISTORY_SHEET)
    ws.append([h for _, h in TEMPLATE_COLS])
    _style_header(ws, len(TEMPLATE_COLS), [k in ("part_no", "die_process") for k, _ in TEMPLATE_COLS])
    for i in range(1, len(TEMPLATE_COLS) + 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = 20
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _validations(ws, name, cols):
    def add(col_key, values):
        idx = [c[0] for c in cols].index(col_key) + 1
        letter = ws.cell(1, idx).column_letter
        dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
        dv.error, dv.errorTitle = "เลือกจากรายการเท่านั้น", "ค่าไม่ถูกต้อง"
        ws.add_data_validation(dv)
        dv.add(f"{letter}2:{letter}2000")
    keys = [c[0] for c in cols]
    for k in ("active", "is_active", "backup"):
        if k in keys:
            add(k, ["Y", "N"])
    if "role" in keys:
        add("role", [r for r in ROLES if r != "superadmin"])
    if name == "7_Approvers":
        add("step", STEP_CODES)
    if "shift" in keys:
        add("shift", ["DAY", "NIGHT", "A", "B"])


def _company_value(company, key, default):
    s = company.settings
    if key == "company_name":
        return company.name
    if key == "logo_text":
        return company.logo_text
    if key == "dr_prefix":
        return company.dr_prefix
    if key == "dr_style":
        return company.dr_style
    if key == "dr_start_no":
        return company.dr_seq or 0
    if key == "pm_skip_foreman":
        return "Y" if s.get("pm_skip_foreman") else "N"
    m = re.match(r"(DAY|NIGHT)_(start|end|breaks)$", key)
    if m:
        sh = (s.get("shifts") or {}).get(m.group(1)) or {}
        if m.group(2) == "breaks":
            return ", ".join(f"{a}-{b}" for a, b in sh.get("breaks", []))
        return sh.get(m.group(2), default)
    return s.get(key, default)


def _current_rows(db, company, name):
    cid = company.id
    yn = lambda b: "Y" if b else "N"
    if name == "2_Lines":
        return [[l.name, yn(l.is_active)] for l in db.query(ProdLine).filter_by(company_id=cid).order_by(ProdLine.name)]
    if name == "3_Customers":
        return [[c.name, yn(c.is_active)] for c in db.query(Customer).filter_by(company_id=cid).order_by(Customer.name)]
    if name == "4_Problems":
        return [[p.name_th, p.name_en, p.sort] for p in db.query(Problem).filter_by(company_id=cid).order_by(Problem.sort)]
    if name == "5_RootCauses":
        return [[p.name_th, p.name_en, p.sort] for p in db.query(RootCause).filter_by(company_id=cid).order_by(RootCause.sort)]
    if name == "6_Users":
        return [[u.username, u.full_name, u.role, "", u.emp_code, u.position, u.department, u.line, u.shift, u.email, u.phone, u.skills, yn(u.is_active)]
                for u in db.query(User).filter(User.company_id == cid, User.password_hash != "!").order_by(User.role, User.full_name)]
    if name == "7_Approvers":
        return [[r.step, r.line or "", r.user.username if r.user else "", yn(r.is_backup)]
                for r in db.query(ApproverRule).filter_by(company_id=cid).order_by(ApproverRule.step, ApproverRule.line)]
    if name == "8_Dies":
        return [[d.customer, d.model, d.part_no, d.part_name, d.die_process, d.die_code, d.line, d.machine_no, d.pm_interval_days or 0,
                 d.pm_interval_strokes or 0, d.stroke_count or 0, d.pm_last_date.strftime("%d-%m-%Y") if d.pm_last_date else "", yn(d.is_active)]
                for d in db.query(Die).filter_by(company_id=cid).order_by(Die.customer, Die.part_no, Die.die_process)]
    return None


# ================================================================== import
def _sheet_rows(wb, name, cols):
    """Rows as dicts keyed by column key. Matches headers by position and by text (tolerates re-ordered columns)."""
    ws = wb[name]
    raw = list(ws.iter_rows(values_only=True))
    if not raw:
        return []
    heads = [clean(h).lower() if h is not None else "" for h in raw[0]]
    pos = {}
    for i, c in enumerate(cols):
        names = {c[1].lower(), c[0].lower(), c[1].lower().replace(" *", "")}
        idx = next((j for j, h in enumerate(heads) if h in names or h.replace(" *", "") in names), None)
        pos[c[0]] = idx if idx is not None else (i if i < len(heads) else None)
    out = []
    for n, r in enumerate(raw[1:], 2):
        if not any(v not in (None, "") for v in r):
            continue
        d = {k: (r[i] if i is not None and i < len(r) else None) for k, i in pos.items()}
        d["_row"] = n
        out.append(d)
    return out


def _s(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return clean(str(v)) if not isinstance(v, (datetime, date)) else v


def _yes(v, default=True):
    s = _s(v)
    return default if s == "" else s.upper() in YES


def _int(v, default=0):
    try:
        return int(float(_s(v) or default))
    except (TypeError, ValueError):
        return None


def _date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = _s(v)
    for f in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


HM = re.compile(r"^([01]?\d|2[0-3]):[0-5]\d$")


def _breaks(text):
    out = []
    for part in re.split(r"[,;\n]+", _s(text)):
        part = part.strip()
        if not part:
            continue
        m = re.match(r"^(\d{1,2}:\d{2})\s*[-–]\s*(\d{1,2}:\d{2})$", part)
        if not m or not HM.match(m.group(1)) or not HM.match(m.group(2)):
            raise ValueError(f"เวลาพัก '{part}' ต้องเป็น HH:MM-HH:MM")
        out.append([m.group(1).zfill(5), m.group(2).zfill(5)])
    return out


def process_kit(db, company, data: bytes, importer, commit=False) -> dict:
    """Validate (commit=False) or import (commit=True) the whole kit. Returns a report per sheet."""
    from ..auth import hash_password
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception as e:
        return {"fatal": f"เปิดไฟล์ไม่ได้ — ต้องเป็น .xlsx ที่ดาวน์โหลดจากระบบ ({e})", "sheets": [], "ok": False}
    cid = company.id
    report = []

    def sec(name, title):
        r = {"sheet": name, "title": title, "rows": 0, "created": 0, "updated": 0, "errors": [], "warnings": [], "present": name in wb.sheetnames}
        report.append(r)
        return r

    sp = db.begin_nested()
    try:
        # ---------- 1 company
        r = sec("1_Company", "ตั้งค่าบริษัท")
        if r["present"]:
            vals = {}
            for row in wb["1_Company"].iter_rows(min_row=2, values_only=True):
                if row and row[0]:
                    vals[_s(row[0])] = row[2] if len(row) > 2 else None
            r["rows"] = len(vals)
            c = db.get(Company, cid)
            upd = {}
            shifts = dict(c.settings.get("shifts") or {})
            for key, val in vals.items():
                sv = _s(val)
                try:
                    if key == "company_name" and sv:
                        c.name = sv
                    elif key == "logo_text" and sv:
                        c.logo_text = sv[:60]
                    elif key == "dr_prefix" and sv:
                        c.dr_prefix = sv[:16]
                    elif key == "dr_style" and sv:
                        if sv not in ("running", "monthly"):
                            raise ValueError("dr_style ต้องเป็น running หรือ monthly")
                        c.dr_style = sv
                    elif key == "dr_start_no" and sv:
                        n = _int(sv)
                        if n is None or n < 0:
                            raise ValueError("dr_start_no ต้องเป็นตัวเลข")
                        if n > (c.dr_seq or 0):
                            c.dr_seq = n
                    elif key == "emergency_required_day" and sv:
                        if sv not in ("<1", "3", "5", ">7"):
                            raise ValueError("emergency_required_day ต้องเป็น <1 / 3 / 5 / >7")
                        upd[key] = sv
                    elif key == "approval_reminder_hours" and sv:
                        n = _int(sv)
                        if n is None:
                            raise ValueError("approval_reminder_hours ต้องเป็นตัวเลข")
                        upd[key] = n
                    elif key == "pm_skip_foreman" and sv:
                        upd[key] = _yes(sv, False)
                    elif key == "approver_visibility" and sv:
                        if sv not in ("own", "all"):
                            raise ValueError("approver_visibility ต้องเป็น own หรือ all")
                        upd[key] = sv
                    elif key in ("tv_working_hours", "spare_app_url"):
                        upd[key] = sv
                    else:
                        m = re.match(r"(DAY|NIGHT)_(start|end|breaks)$", key)
                        if m and sv:
                            sh = dict(shifts.get(m.group(1)) or {"start": "08:00", "end": "20:00", "breaks": []})
                            if m.group(2) == "breaks":
                                sh["breaks"] = _breaks(sv)
                            else:
                                if not HM.match(sv):
                                    raise ValueError(f"{key} ต้องเป็น HH:MM")
                                sh[m.group(2)] = sv.zfill(5)
                            shifts[m.group(1)] = sh
                except ValueError as e:
                    r["errors"].append(f"{key}: {e}")
            upd["shifts"] = shifts
            upd["setup_company_done"] = True
            c.set_settings(upd)
            r["updated"] = len(vals)

        # ---------- simple masters
        def master(name, model, key_field, fields):
            r = sec(name, SHEETS[name][0])
            if not r["present"]:
                return set()
            rows = _sheet_rows(wb, name, SHEETS[name][1])
            r["rows"] = len(rows)
            have = {getattr(o, key_field).strip().lower(): o for o in db.query(model).filter_by(company_id=cid)}
            seen = set()
            for row in rows:
                val = _s(row[fields[0]])
                if name == "2_Lines":
                    val = norm_line(val)
                if not val:
                    r["errors"].append(f"แถว {row['_row']}: ช่อง {SHEETS[name][1][0][1]} ว่าง")
                    continue
                k = val.lower()
                if k in seen:
                    r["warnings"].append(f"แถว {row['_row']}: '{val}' ซ้ำในไฟล์ (ใช้แถวแรก)")
                    continue
                seen.add(k)
                obj = have.get(k)
                if obj is None:
                    obj = model(company_id=cid)
                    setattr(obj, key_field, val[:120])
                    db.add(obj)
                    have[k] = obj
                    r["created"] += 1
                else:
                    r["updated"] += 1
                if "active" in row:
                    obj.is_active = _yes(row.get("active"))
                if "name_en" in row:
                    obj.name_en = (_s(row.get("name_en")) or val)[:120]
                    s = _int(row.get("sort"), 0)
                    obj.sort = s if s is not None else 0
            db.flush()
            return seen

        lines = master("2_Lines", ProdLine, "name", ["line"])
        master("3_Customers", Customer, "name", ["customer"])
        master("4_Problems", Problem, "name_th", ["name_th"])
        master("5_RootCauses", RootCause, "name_th", ["name_th"])
        all_lines = {l.name.lower() for l in db.query(ProdLine).filter_by(company_id=cid)}

        # ---------- users
        r = sec("6_Users", "ผู้ใช้")
        if r["present"]:
            rows = _sheet_rows(wb, "6_Users", SHEETS["6_Users"][1])
            r["rows"] = len(rows)
            res = upsert_users(db, cid, rows, hash_password, all_lines)
            r.update(created=res["created"], updated=res["updated"])
            r["errors"] += res["errors"]
            r["warnings"] += res["warnings"]

        # ---------- approvers
        r = sec("7_Approvers", "ผังผู้อนุมัติ")
        if r["present"]:
            rows = _sheet_rows(wb, "7_Approvers", SHEETS["7_Approvers"][1])
            r["rows"] = len(rows)
            users = {u.username: u for u in db.query(User).filter(User.company_id == cid)}
            have = {(x.step, (x.line or "").lower(), x.user_id): x for x in db.query(ApproverRule).filter_by(company_id=cid)}
            for row in rows:
                step = _s(row["step"]).upper()
                line = norm_line(_s(row["line"]))
                uname = _s(row["username"]).lower()
                if step not in STEP_CODES:
                    r["errors"].append(f"แถว {row['_row']}: step '{step}' ต้องเป็น {' / '.join(STEP_CODES)}")
                    continue
                u = users.get(uname)
                if not u:
                    r["errors"].append(f"แถว {row['_row']}: ไม่พบ username '{uname}' (เพิ่มในชีต 6_Users ก่อน)")
                    continue
                if not u.is_active:
                    r["warnings"].append(f"แถว {row['_row']}: {uname} ถูกระงับ — จะไม่ได้รับงาน")
                if line and line.lower() not in all_lines:
                    r["warnings"].append(f"แถว {row['_row']}: Line '{line}' ไม่อยู่ในชีต 2_Lines")
                key = (step, line.lower(), u.id)
                rule = have.get(key)
                if rule is None:
                    rule = ApproverRule(company_id=cid, step=step, line=line, user_id=u.id)
                    db.add(rule)
                    have[key] = rule
                    r["created"] += 1
                else:
                    r["updated"] += 1
                rule.is_backup = _yes(row.get("backup"), False)
            db.flush()

        # ---------- dies
        r = sec("8_Dies", "ทะเบียนแม่พิมพ์")
        if r["present"]:
            rows = _sheet_rows(wb, "8_Dies", SHEETS["8_Dies"][1])
            r["rows"] = len(rows)
            dies = {(d.part_no.lower(), (d.die_process or "").lower()): d for d in db.query(Die).filter_by(company_id=cid)}
            custs = {c.name.lower() for c in db.query(Customer).filter_by(company_id=cid)}
            seen = set()
            for row in rows:
                pn, proc = _s(row["part_no"]), _s(row["die_process"])
                if not pn or not proc:
                    r["errors"].append(f"แถว {row['_row']}: ต้องมี PartNo และ DieProcess")
                    continue
                k = (pn.lower(), proc.lower())
                if k in seen:
                    r["warnings"].append(f"แถว {row['_row']}: {pn} / {proc} ซ้ำในไฟล์")
                    continue
                seen.add(k)
                d = dies.get(k)
                if d is None:
                    d = Die(company_id=cid, part_no=pn[:80], die_process=proc[:120])
                    db.add(d)
                    dies[k] = d
                    r["created"] += 1
                else:
                    r["updated"] += 1
                line = norm_line(_s(row["line"]))
                cust = _s(row["customer"])
                for f, v in (("customer", cust), ("model", _s(row["model"])), ("part_name", _s(row["part_name"])),
                             ("die_code", _s(row["die_code"])), ("line", line), ("machine_no", _s(row["machine_no"]))):
                    if v:
                        setattr(d, f, v)
                pmi = _int(row["pm_interval_days"], 0)
                if pmi is None:
                    r["errors"].append(f"แถว {row['_row']}: PM Interval Days ต้องเป็นตัวเลข")
                else:
                    d.pm_interval_days = pmi
                for key, label in (("pm_interval_strokes", "PM Interval Strokes"), ("stroke_count", "Stroke Count")):
                    if _s(row.get(key)):
                        v = _int(str(row.get(key)).replace(",", ""), 0)
                        if v is None or v < 0:
                            r["errors"].append(f"แถว {row['_row']}: {label} ต้องเป็นตัวเลข")
                        else:
                            setattr(d, key, v)
                if _s(row["pm_last_date"]):
                    dt = _date(row["pm_last_date"])
                    if not dt:
                        r["errors"].append(f"แถว {row['_row']}: PM Last Date ต้องเป็น DD-MM-YYYY")
                    else:
                        d.pm_last_date = dt
                        if d.pm_interval_days:
                            from datetime import timedelta
                            d.pm_next_date = dt + timedelta(days=d.pm_interval_days)
                d.is_active = _yes(row.get("active"))
                if line and line.lower() not in all_lines:
                    db.add(ProdLine(company_id=cid, name=line))
                    all_lines.add(line.lower())
                    r["warnings"].append(f"แถว {row['_row']}: เพิ่ม Line '{line}' ให้อัตโนมัติ (ไม่มีในชีต 2_Lines)")
                if cust and cust.lower() not in custs:
                    db.add(Customer(company_id=cid, name=cust))
                    custs.add(cust.lower())
                db.flush()

        # ---------- history (validated here, imported by the proven history importer)
        r = sec(HISTORY_SHEET, "ใบงานเก่า (ไม่บังคับ)")
        hist_rows = []
        if r["present"]:
            ws = wb[HISTORY_SHEET]
            raw = list(ws.iter_rows(values_only=True))
            if raw:
                heads = [clean(h) for h in raw[0]]
                hist_rows = [dict(zip(heads, x)) for x in raw[1:] if any(v not in (None, "") for v in x)]
            r["rows"] = len(hist_rows)
            for i, h in enumerate(hist_rows, 2):
                if not clean(h.get("Part No")) or not clean(h.get("Die Process")):
                    r["errors"].append(f"แถว {i}: ต้องมี Part No และ Die Process")
            if len(r["errors"]) > 30:
                r["errors"] = r["errors"][:30] + [f"... และอีก {len(r['errors']) - 30} แถว"]
    except Exception as e:
        sp.rollback()
        db.rollback()
        return {"fatal": f"ไฟล์ผิดรูปแบบ: {e}", "sheets": report, "ok": False}

    errors = sum(len(r["errors"]) for r in report)
    if commit and not errors:
        sp.commit()
        db.commit()
        if hist_rows:
            from .import_data import import_requests
            res = import_requests(db, db.get(Company, cid), "kit.xlsx", b"", importer, rows=hist_rows)
            h = next(r for r in report if r["sheet"] == HISTORY_SHEET)
            h["created"], h["updated"] = res["created"], 0
            h["warnings"] += [f"ข้าม {res['skipped']} แถว (Legacy ID ซ้ำ/ไม่มี Part No)"] if res["skipped"] else []
            h["errors"] += res["errors"]
        return {"sheets": report, "ok": True, "imported": True, "errors": errors}
    sp.rollback()
    db.rollback()
    return {"sheets": report, "ok": errors == 0, "imported": False, "errors": errors}


def upsert_users(db, cid, rows, hash_password, lines=None) -> dict:
    """Shared by the start-up kit and Admin › ผู้ใช้ › นำเข้า. rows: dicts with USER keys + optional _row."""
    created = updated = 0
    errors, warnings = [], []
    seen = set()
    for i, r in enumerate(rows, 2):
        n = r.get("_row", i)
        g = lambda k: _s(r.get(k))
        uname = g("username").lower()
        if not uname and not g("full_name"):
            continue
        if not uname or not g("full_name"):
            errors.append(f"แถว {n}: ต้องมี username และ full_name")
            continue
        if not re.match(r"^[a-z0-9._@-]{3,64}$", uname):
            errors.append(f"แถว {n}: username '{uname}' ใช้ได้เฉพาะ a-z 0-9 . _ - (3-64 ตัว)")
            continue
        if uname in seen:
            errors.append(f"แถว {n}: username '{uname}' ซ้ำในไฟล์")
            continue
        seen.add(uname)
        role = g("role") or "production"
        if role not in ROLES or role == "superadmin":
            errors.append(f"แถว {n}: role '{role}' ไม่ถูกต้อง")
            continue
        email = g("email")
        if email and not re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", email):
            errors.append(f"แถว {n}: email '{email}' ไม่ถูกต้อง")
            continue
        u = db.query(User).filter(User.username == uname).first()
        if u and u.company_id != cid:
            errors.append(f"แถว {n}: username {uname} ถูกใช้โดยบริษัทอื่น")
            continue
        pw = g("password")
        if not u:
            if not pw:
                errors.append(f"แถว {n}: ผู้ใช้ใหม่ {uname} ต้องมี password")
                continue
            # an inactive placeholder created by the history import with the same name becomes this account
            u = db.query(User).filter(User.company_id == cid, User.full_name == g("full_name"), User.password_hash == "!").first()
            if u:
                u.username = uname
                updated += 1
            else:
                u = User(company_id=cid, username=uname, full_name=g("full_name"), password_hash="!")
                db.add(u)
                created += 1
            u.password_hash, u.must_change_pw = hash_password(pw), True
        else:
            updated += 1
            if pw:
                u.password_hash, u.must_change_pw = hash_password(pw), True
        if len(pw) and len(pw) < 6:
            warnings.append(f"แถว {n}: รหัสผ่านของ {uname} สั้นกว่า 6 ตัว")
        u.full_name, u.role = g("full_name")[:120], role
        for k in ("emp_code", "position", "department", "line", "shift", "email", "phone", "skills"):
            if k in r:
                v = g(k)
                if k == "line":
                    v = norm_line(v)
                    if v and lines is not None and v.lower() not in lines:
                        warnings.append(f"แถว {n}: Line '{v}' ของ {uname} ไม่มีใน Master Line")
                setattr(u, k, v or None)
        u.is_active = _yes(r.get("is_active"))
        if role in ("foreman", "die_leader", "qc") and not g("email"):
            warnings.append(f"แถว {n}: {uname} ({role}) ไม่มีอีเมล — จะไม่ได้รับอีเมลแจ้งอนุมัติ")
    db.flush()
    return {"created": created, "updated": updated, "errors": errors, "warnings": warnings}


# ================================================================== readiness
def readiness(db, company) -> dict:
    """Go-live checklist. level: req (must pass) / rec (recommended)."""
    from .. import approvals
    from ..config import settings
    cid = company.id
    s = company.settings
    items = []

    def add(group, level, ok, title, detail, fix, fix_label="แก้ไข"):
        items.append({"group": group, "level": level, "ok": bool(ok), "title": title, "detail": detail, "fix": fix, "fix_label": fix_label})

    n = lambda model, **kw: db.query(model).filter_by(company_id=cid, **kw).count()
    users = db.query(User).filter(User.company_id == cid, User.is_active == True).all()
    by_role = {}
    for u in users:
        by_role.setdefault(u.role, []).append(u)

    add("1. ตั้งค่าบริษัท", "rec", s.get("setup_company_done") or company.dr_prefix not in ("", "DR"),
        "รหัสเลขใบงาน / เวลากะ / เวลาพัก", f"Prefix: {company.dr_prefix} · เลขล่าสุด {company.dr_seq or 0}", "/admin/settings")
    breaks_ok = all((s.get("shifts") or {}).get(k, {}).get("breaks") for k in SHIFT_KEYS)
    add("1. ตั้งค่าบริษัท", "req", breaks_ok, "เวลาพักกะเช้า/กะดึก (ใช้หักเวลาซ่อมสุทธิ)", "ครบ" if breaks_ok else "ยังไม่ได้กำหนดเวลาพัก", "/admin/settings")

    nl, nc, npb, nrc = n(ProdLine, is_active=True), n(Customer, is_active=True), n(Problem, is_active=True), n(RootCause, is_active=True)
    add("2. Master data", "req", nl > 0, "Line การผลิต", f"{nl} Line", "/admin/masters")
    add("2. Master data", "rec", nc > 0, "ลูกค้า", f"{nc} ราย", "/admin/masters")
    add("2. Master data", "req", npb > 0, "รายการปัญหา (Production เลือก)", f"{npb} รายการ", "/admin/masters")
    add("2. Master data", "req", nrc > 0, "รายการสาเหตุ (ช่างเลือก)", f"{nrc} รายการ", "/admin/masters")

    nd = n(Die, is_active=True)
    add("3. ทะเบียนแม่พิมพ์", "req", nd > 0, "มีแม่พิมพ์ในทะเบียน", f"{nd} แม่พิมพ์", "/admin/setup#kit", "นำเข้า")
    if nd:
        no_line = db.query(Die).filter(Die.company_id == cid, Die.is_active == True, (Die.line == None) | (Die.line == "")).count()
        no_cust = db.query(Die).filter(Die.company_id == cid, Die.is_active == True, (Die.customer == None) | (Die.customer == "")).count()
        add("3. ทะเบียนแม่พิมพ์", "rec", no_line == 0, "ทุกแม่พิมพ์ระบุ Line", f"ไม่มี Line {no_line} แม่พิมพ์" if no_line else "ครบ", "/dies")
        add("3. ทะเบียนแม่พิมพ์", "rec", no_cust == 0, "ทุกแม่พิมพ์ระบุลูกค้า", f"ไม่มีลูกค้า {no_cust} แม่พิมพ์" if no_cust else "ครบ", "/dies")
        from sqlalchemy import or_ as _or
        pm = db.query(Die).filter(Die.company_id == cid, Die.is_active == True, _or(Die.pm_interval_days > 0, Die.pm_interval_strokes > 0)).count()
        add("3. ทะเบียนแม่พิมพ์", "rec", pm >= max(1, nd // 10), "กำหนดรอบ PM (อย่างน้อย 10% ของแม่พิมพ์)", f"มีแผน PM {pm} / {nd} แม่พิมพ์", "/pm")

    role_th = {"production": "Production (ผู้แจ้ง)", "foreman": "Foreman / Engineer", "die_leader": "หัวหน้า Die Maintenance",
               "technician": "ช่างซ่อมแม่พิมพ์", "qc": "QC"}
    for role, label in role_th.items():
        k = len(by_role.get(role, []))
        add("4. ผู้ใช้", "req", k > 0, f"มีผู้ใช้บทบาท {label}", f"{k} คน", "/admin/users")
    add("4. ผู้ใช้", "rec", len(by_role.get("admin", [])) > 0, "มีผู้ดูแลระบบของบริษัท (admin)", f"{len(by_role.get('admin', []))} คน", "/admin/users")
    pending = sum(1 for u in users if u.must_change_pw and not u.last_login)
    add("4. ผู้ใช้", "rec", pending == 0, "ทุกคนเข้าระบบและตั้งรหัสผ่านของตัวเองแล้ว",
        f"ยังไม่เคยเข้าระบบ {pending} คน" if pending else "ครบ", "/admin/users", "ดูรายชื่อ")
    techs = by_role.get("technician", [])
    nophoto = sum(1 for u in techs if not u.photo_path)
    add("4. ผู้ใช้", "rec", techs and nophoto == 0, "รูปพนักงานทีมช่าง (แสดงบนจอ TV)",
        f"ยังไม่มีรูป {nophoto} / {len(techs)} คน" if nophoto else "ครบ", "/admin/users#photos", "อัปโหลดรูป")

    approvals.invalidate(db)
    rows, warns = approvals.coverage(db, cid)
    empty = [w for w in warns if "ไม่มีผู้อนุมัติ" in w]
    nomail = [w for w in warns if "อีเมล" in w]
    nrules = n(ApproverRule)
    add("5. ผังผู้อนุมัติ", "req", not empty, "ทุก Line มีผู้อนุมัติครบทุกขั้น", "; ".join(empty[:4]) or f"ครบ ({len(rows)} Line)", "/admin/approvers")
    add("5. ผังผู้อนุมัติ", "rec", nrules > 0, "กำหนดผังผู้อนุมัติของบริษัทเอง (ไม่ใช้ค่าอัตโนมัติ)",
        f"{nrules} กฎ" if nrules else "ยังใช้ 'บทบาท + Line ประจำ' อัตโนมัติ", "/admin/approvers")
    add("5. ผังผู้อนุมัติ", "rec", not nomail, "ผู้อนุมัติมีอีเมล", "; ".join(nomail[:3]) or "ครบ", "/admin/users")

    from .notify import email_config
    ec = email_config(company, db)
    nch = n(NotificationChannel, is_active=True)
    add("6. แจ้งเตือน & Backup", "rec", ec["enabled"], "อีเมลถึงผู้อนุมัติรายบุคคล (Brevo)", "พร้อม" if ec["enabled"] else "ยังไม่ได้ใส่ Brevo API key / อีเมลผู้ส่ง", "/admin/channels")
    add("6. แจ้งเตือน & Backup", "rec", nch > 0, "ช่องทางแจ้งเตือนกลุ่ม (LINE / Telegram / Discord)", f"{nch} ช่องทาง", "/admin/channels")
    from . import images
    offsite = images.using_s3() or db.query(NotificationChannel).filter(NotificationChannel.company_id == cid, NotificationChannel.is_active == True,
                                                                        NotificationChannel.events_json.like('%backup%')).count() > 0
    add("6. แจ้งเตือน & Backup", "rec", s.get("auto_backup") and offsite, "Backup อัตโนมัติส่งออกนอกเซิร์ฟเวอร์",
        "เปิดอยู่" if (s.get("auto_backup") and offsite) else "ยังไม่มีปลายทาง (Bucket หรือช่องทางที่ติ๊ก backup)", "/admin/backup")
    add("6. แจ้งเตือน & Backup", "rec", not settings.BASE_URL.startswith("http://localhost") and not settings.BASE_URL.startswith("http://127."),
        "URL ของระบบ (ใช้ในลิงก์อีเมล/QR)", settings.BASE_URL, "/admin/system")

    req = [i for i in items if i["level"] == "req"]
    rec = [i for i in items if i["level"] == "rec"]
    req_ok, rec_ok = sum(i["ok"] for i in req), sum(i["ok"] for i in rec)
    groups = {}
    for i in items:
        groups.setdefault(i["group"], []).append(i)
    return {"items": items, "groups": groups, "req_total": len(req), "req_ok": req_ok, "rec_total": len(rec), "rec_ok": rec_ok,
            "ready": req_ok == len(req), "pct": round(100 * (req_ok + rec_ok) / max(len(items), 1))}
