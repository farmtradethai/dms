"""Historic data import (Power Apps export or the Excel template) + die register import.
Everything is cleaned on the way in: trimmed strings, Line/Customer spelling normalised,
US dates parsed, and the "20.00-22.00" time range technicians type at the end of HowToFix
is converted to actual repair minutes so MTTR works for old data too."""
import csv
import io
import re
from datetime import date, datetime, timedelta

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy.orm import Session

from ..database import now_local
from ..models import (ST, Assignment, Die, Problem, RepairRequest, RootCause, StatusHistory, User, Customer, ProdLine)

TEMPLATE_COLS = [
    ("legacy_id", "Legacy ID (AHMDxxxxx)"), ("request_type", "Type (REPAIR/PM/ACCIDENT)"),
    ("request_date", "Request Date (DD-MM-YYYY HH:MM)"), ("requester", "Requester"), ("phone", "Phone"),
    ("line", "Line"), ("customer", "Customer"), ("model", "Model"), ("part_no", "Part No"), ("part_name", "Part Name"),
    ("die_process", "Die Process"), ("required_day", "Required Day (<1/3/5/>7)"), ("stock_oh", "Stock OH"),
    ("problem", "Problem"), ("other_details", "Other Details"), ("status", "Status (CLOSED/WAIT_QC/...)"),
    ("foreman", "Foreman Name"), ("foreman_date", "Foreman Approve Date"), ("leader", "Die Leader Name"),
    ("leader_date", "Leader Approve Date"), ("technicians", "Technicians (comma separated)"),
    ("received_date", "Tech Received Date"), ("complete_date", "Complete Date"), ("root_cause", "Root Cause"),
    ("action", "Action / How to fix"), ("repair_minutes", "Net Repair Minutes"), ("tryout_date", "Try-out Accept Date"),
    ("tryout_by", "Try-out By"), ("qc_date", "QC Check Date"), ("qc_by", "QC By"), ("qc_result", "QC Result (OK/NG)"),
    ("qc_comment", "QC Comment"), ("downtime_minutes", "Downtime Minutes"), ("parts_cost", "Parts Cost"),
]

# Power Apps column -> template key
POWERAPPS_MAP = {
    "DIEId": "legacy_id", "RepairType": "request_type", "RequestDate": "request_date", "UserFullName": "requester",
    "UserPhone": "phone", "Line": "line", "CustomerName": "customer", "Model": "model", "PartNumber": "part_no",
    "PartName": "part_name", "DieProcess": "die_process", "RequiredDAY": "required_day", "StockOH": "stock_oh",
    "ProblemDetails": "problem", "OtherDetails": "other_details", "DIEStatus": "status", "FApproveName": "foreman",
    "FApproveDate": "foreman_date", "DIELeaderApproveName": "leader", "DIELeaderApproveDate": "leader_date",
    "TechnicianName": "technicians", "TechReceiveJobDate": "received_date", "CompleteFixDate": "complete_date",
    "RootCause": "root_cause", "HowToFix": "action", "LineLDAcceptDate": "tryout_date", "LineLDAcceptName": "tryout_by",
    "QCCheckDate": "qc_date", "QCCheckName": "qc_by", "QCResult": "qc_result", "QCCheckComment": "qc_comment",
}
STATUS_MAP = {
    "closed request": ST["CLOSED"], "waiting line leader accepted": ST["REPAIR_DONE"], "waiting qc check": ST["WAIT_QC"],
    "die tech received job": ST["IN_PROGRESS"], "waiting die tech received job": ST["ASSIGNED"],
    "waiting die mt approve": ST["WAIT_LEADER"], "waiting foreman approve": ST["WAIT_FOREMAN"],
    "die mt rejected": ST["LEADER_REJECTED"], "foreman rejected": ST["FOREMAN_REJECTED"], "qc rejected": ST["NG_RETURN"],
    "line leader rejected": ST["RETURNED"], "repair updated progress": ST["IN_PROGRESS"], "die repair done": ST["REPAIR_DONE"],
    "waiting close request": ST["WAIT_QC"], "new request repair": ST["WAIT_FOREMAN"],
}
TYPE_MAP = {"ซ่อมแม่พิมพ์": "REPAIR", "pm": "PM", "die accident": "ACCIDENT", "repair": "REPAIR", "accident": "ACCIDENT"}
DATE_FORMATS = ["%m/%d/%Y %I:%M %p", "%m/%d/%Y %H:%M", "%d-%m-%Y %H:%M", "%d/%m/%Y %H:%M", "%Y-%m-%d %H:%M:%S",
                "%Y-%m-%d %H:%M", "%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%m/%d/%Y"]
TIME_RANGE = re.compile(r"(\d{1,2})[.:](\d{2})\s*[-–—*;]\s*(\d{1,2})[.:](\d{2})")


def clean(s):
    if s is None:
        return ""
    if isinstance(s, (datetime, date)):
        return s
    return re.sub(r"\s+", " ", str(s)).strip()


def norm_line(s: str) -> str:
    s = clean(s)
    return re.sub(r"^l[i1]ne\s*", "Line", s, flags=re.I).replace("Line ", "Line").replace("LineIMV", "Line IMV") if s else s


US_FIRST = DATE_FORMATS
DMY_FIRST = ["%d-%m-%Y %H:%M", "%d/%m/%Y %H:%M", "%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d",
             "%m/%d/%Y %I:%M %p", "%m/%d/%Y %H:%M", "%m/%d/%Y"]


def parse_dt(v, prefer_us=True):
    """Power Apps CSV exports are US (MM/DD/YYYY, usually with AM/PM); Thai Excel files are DD/MM/YYYY.
    prefer_us=False reads day first so 05/03/2025 = 5 March, not May 3."""
    if not v:
        return None
    if isinstance(v, datetime):
        return v
    if isinstance(v, date):
        return datetime.combine(v, datetime.min.time())
    s = clean(v)
    if re.search(r"\b(AM|PM)\b", s, re.I):
        prefer_us = True                      # AM/PM only appears in the Power Apps (US) export
    for f in (US_FIRST if prefer_us else DMY_FIRST):
        try:
            return datetime.strptime(s, f)
        except ValueError:
            continue
    return None


def minutes_from_text(text: str):
    """'เปลี่ยน insert ใหม่ 20.00-22.00' -> 120 ; crosses midnight handled."""
    if not text:
        return None
    m = TIME_RANGE.search(str(text))
    if not m:
        return None
    h1, m1, h2, m2 = (int(x) for x in m.groups())
    if h1 > 24 or h2 > 24 or m1 > 59 or m2 > 59:
        return None
    a, b = h1 * 60 + m1, h2 * 60 + m2
    if b < a:
        b += 24 * 60
    return b - a if 0 < b - a <= 24 * 60 else None


def required_day_norm(v: str) -> str:
    s = clean(v).replace(" ", "").lower()
    if s.startswith("<1") or s in ("1", "1day", "emergency", "ด่วน"):
        return "<1"
    if s.startswith("3"):
        return "3"
    if s.startswith("5"):
        return "5"
    if s.startswith(">7") or s.startswith("7"):
        return ">7"
    return "3"


def build_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Requests"
    ws.append([h for _, h in TEMPLATE_COLS])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F5FBF")
    ws.append(["AHMD02072", "REPAIR", "27-02-2025 15:49", "AH Press", "0804524759", "Line3", "TOYOTA", "640A",
               "44591-0K110", "BRACKET BRAKE ACTUATOR NO.1", "OP10 Progressive", "3", 1300, "อื่นๆ",
               "เปลี่ยน Little ประคองงาน", "CLOSED", "Surapin Takengsook", "27-02-2025 15:50", "Panat Akshee",
               "28-02-2025 10:17", "นาย ณภัทร ปารสูตร", "28-02-2025 13:24", "04-03-2025 09:34", "อื่นๆ",
               "เปลี่ยนอะไหล่ประครองเหล็กใหม่ 20.00-22.00", 120, "05-03-2025 08:00", "AH Press", "06-03-2025 10:00",
               "Tanakrit Intarasung", "OK", "QC OK", 0, 0])
    for i in range(1, len(TEMPLATE_COLS) + 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = 22
    info = wb.create_sheet("วิธีใช้")
    for line in ["1. กรอกข้อมูลเก่าในชีต Requests (บรรทัดที่ 2 เป็นตัวอย่าง ลบทิ้งได้)",
                 "2. วันที่ใช้รูปแบบ DD-MM-YYYY HH:MM (หรือไฟล์ Export จาก Power Apps ใช้ MM/DD/YYYY ก็อ่านได้)",
                 "3. Status ใช้รหัส: CLOSED, WAIT_QC, REPAIR_DONE, IN_PROGRESS, ASSIGNED, WAIT_LEADER, WAIT_FOREMAN หรือชื่อสถานะ Power Apps เดิม",
                 "4. ถ้า Net Repair Minutes ว่าง ระบบจะอ่านช่วงเวลาจากช่อง Action เช่น '... 20.00-22.00' อัตโนมัติ",
                 "5. อัปโหลดไฟล์นี้ หรือไฟล์ CSV Export จาก Power Apps (DieRepairList.csv) ได้โดยตรงที่ Admin > นำเข้าข้อมูลเก่า",
                 "6. รายการที่ Legacy ID ซ้ำกับที่มีอยู่แล้วจะถูกข้าม"]:
        info.append([line])
    info.column_dimensions["A"].width = 120
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def read_rows(filename: str, data: bytes):
    if filename.lower().endswith(".csv"):
        text = data.decode("utf-8-sig", errors="replace")
        return list(csv.DictReader(io.StringIO(text)))
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb["Requests"] if "Requests" in wb.sheetnames else wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [clean(h) for h in rows[0]]
    return [dict(zip(headers, r)) for r in rows[1:] if any(v not in (None, "") for v in r)]


def _normalise_record(raw: dict) -> dict:
    rec = {}
    header_to_key = {h: k for k, h in TEMPLATE_COLS}
    for k, v in raw.items():
        key = header_to_key.get(k) or POWERAPPS_MAP.get(k)
        if key:
            rec[key] = v
    return rec


def _find_user(db, company_id, name):
    name = clean(name)
    if not name:
        return None
    return db.query(User).filter(User.company_id == company_id, User.full_name == name).first()


def import_requests(db: Session, company, filename: str, data: bytes, importer: User, **kwargs) -> dict:
    rows = kwargs.get("rows") if kwargs.get("rows") is not None else read_rows(filename, data)
    settings = company.settings
    existing = {r.legacy_id for r in db.query(RepairRequest.legacy_id).filter(RepairRequest.company_id == company.id) if r.legacy_id}
    problems = {p.name_th: p for p in db.query(Problem).filter(Problem.company_id == company.id)}
    causes = {c.name_th: c for c in db.query(RootCause).filter(RootCause.company_id == company.id)}
    customers = {c.name for c in db.query(Customer).filter(Customer.company_id == company.id)}
    lines = {l.name for l in db.query(ProdLine).filter(ProdLine.company_id == company.id)}
    techs = {u.full_name: u for u in db.query(User).filter(User.company_id == company.id)}
    users_by_name = dict(techs)
    dies = {(d.part_no, d.die_process): d for d in db.query(Die).filter(Die.company_id == company.id)}
    created, skipped, errors = 0, 0, []
    progress = kwargs.get("progress")
    from ..routers.requests import next_dr_no  # local import to avoid a cycle

    for i, raw in enumerate(rows, 2):
        sp = db.begin_nested()          # SAVEPOINT: a bad row is undone alone, earlier rows stay
        try:
            rec = _normalise_record(raw)
            us = "DIEId" in raw or "RequestDate" in raw          # Power Apps export -> US dates; our template -> DD-MM-YYYY
            _pdt = lambda v: parse_dt(v, prefer_us=us)
            legacy = clean(rec.get("legacy_id"))
            if legacy and legacy in existing:
                skipped += 1
                sp.commit()
                continue
            part_no, process = clean(rec.get("part_no")), clean(rec.get("die_process"))
            if not part_no:
                skipped += 1
                sp.commit()
                continue
            line = norm_line(rec.get("line"))
            customer = clean(rec.get("customer"))
            if customer and customer not in customers:
                db.add(Customer(company_id=company.id, name=customer)); customers.add(customer)
            if line and line not in lines:
                db.add(ProdLine(company_id=company.id, name=line)); lines.add(line)
            die = dies.get((part_no, process or "-")) or dies.get((part_no, process))
            if not die:
                die = Die(company_id=company.id, part_no=part_no, part_name=clean(rec.get("part_name")),
                          die_process=process or "-", customer=customer, model=clean(rec.get("model")), line=line)
                db.add(die)
                db.flush()
                dies[(part_no, process or "-")] = die
            problem_txt = clean(rec.get("problem"))
            if problem_txt and problem_txt not in problems:
                p = Problem(company_id=company.id, name_th=problem_txt, sort=99)
                db.add(p); db.flush(); problems[problem_txt] = p
            cause_txt = clean(rec.get("root_cause"))
            if cause_txt and cause_txt not in causes:
                c = RootCause(company_id=company.id, name_th=cause_txt, sort=99)
                db.add(c); db.flush(); causes[cause_txt] = c
            status_raw = clean(rec.get("status"))
            status = STATUS_MAP.get(status_raw.lower(), status_raw.upper() if status_raw.upper() in ST.values() else ST["CLOSED"])
            req_date = _pdt(rec.get("request_date")) or now_local()
            rtype = TYPE_MAP.get(clean(rec.get("request_type")).lower(), "REPAIR")
            rd = required_day_norm(rec.get("required_day"))
            action = clean(rec.get("action"))
            mins = rec.get("repair_minutes")
            try:
                mins = int(float(mins)) if mins not in (None, "") else None
            except ValueError:
                mins = None
            if mins is None:
                mins = minutes_from_text(action) or 0
            try:
                stock = int(float(clean(rec.get("stock_oh")) or 0))
            except ValueError:
                stock = None
            days = {"<1": 1, "3": 3, "5": 5, ">7": 7}[rd]
            req = RepairRequest(
                company_id=company.id, dr_no=legacy or next_dr_no(db, company), legacy_id=legacy or None,
                request_type=rtype, die_id=die.id, die_code=die.die_code, customer=customer, model=clean(rec.get("model")),
                part_no=part_no, part_name=clean(rec.get("part_name")) or die.part_name, die_process=process, line=line,
                required_day=rd, is_emergency=(rd == settings.get("emergency_required_day", "<1")),
                target_date=(req_date + timedelta(days=days)).date(), stock_oh=stock,
                problem_id=problems[problem_txt].id if problem_txt else None, problem_text=problem_txt,
                other_details=clean(rec.get("other_details")), requester_name=clean(rec.get("requester")) or "Import",
                requester_phone=clean(rec.get("phone")), department="Production", status=status, created_at=req_date,
                foreman_at=_pdt(rec.get("foreman_date")), leader_at=_pdt(rec.get("leader_date")),
                tech_received_at=_pdt(rec.get("received_date")), tech_completed_at=_pdt(rec.get("complete_date")),
                root_cause_id=causes[cause_txt].id if cause_txt else None, root_cause_text=cause_txt, action_text=action,
                net_repair_minutes=mins, tryout_at=_pdt(rec.get("tryout_date")),
                qc_at=_pdt(rec.get("qc_date")), qc_result=clean(rec.get("qc_result")).upper()[:4] or None,
                qc_comment=clean(rec.get("qc_comment")), progress_pct=100 if _pdt(rec.get("complete_date")) else 0,
            )
            fm = users_by_name.get(clean(rec.get("foreman")))
            if fm:
                req.foreman_id = fm.id
            ld = users_by_name.get(clean(rec.get("leader")))
            if ld:
                req.leader_id = ld.id
            if req.qc_at and status == ST["CLOSED"]:
                req.closed_at = req.qc_at
            if req.tryout_at:
                req.tryout_count = 1
                req.tryout_result = "OK"
            try:
                req.downtime_minutes = int(float(rec.get("downtime_minutes") or 0))
                req.parts_cost = float(rec.get("parts_cost") or 0)
            except (TypeError, ValueError):
                pass
            db.add(req)
            db.flush()
            # technicians (create inactive placeholder users so names stay linkable)
            for name in [clean(x) for x in re.split(r"[,/]", clean(rec.get("technicians")) or "") if clean(x)]:
                u = techs.get(name)
                if not u:
                    u = User(company_id=company.id, username=f"tech_{company.code.lower()}_{len(techs) + 1}_{abs(hash(name)) % 9999}",
                             password_hash="!", full_name=name, role="technician", is_active=False)
                    db.add(u); db.flush(); techs[name] = u; users_by_name[name] = u
                db.add(Assignment(request_id=req.id, user_id=u.id))
            # history reconstructed from the stage timestamps
            stages = [(req.created_at, ST["WAIT_FOREMAN"], "Import"), (req.foreman_at, ST["WAIT_LEADER"], clean(rec.get("foreman"))),
                      (req.leader_at, ST["ASSIGNED"], clean(rec.get("leader"))), (req.tech_received_at, ST["IN_PROGRESS"], clean(rec.get("technicians"))),
                      (req.tech_completed_at, ST["REPAIR_DONE"], clean(rec.get("technicians"))), (req.tryout_at, ST["WAIT_QC"], clean(rec.get("tryout_by"))),
                      (req.qc_at, ST["CLOSED"] if req.qc_result == "OK" else ST["NG_RETURN"], clean(rec.get("qc_by")))]
            prev = None
            for at, st, who in stages:
                if at:
                    db.add(StatusHistory(request_id=req.id, from_status=prev, to_status=st, user_name=who or "Import", at=at,
                                         comment="นำเข้าจากข้อมูลเก่า"))
                    prev = st
            sp.commit()
            created += 1
            if legacy:
                existing.add(legacy)
            if created % 200 == 0:
                db.commit()
                if progress:
                    progress(created, len(rows))
        except Exception as e:
            errors.append(f"row {i}: {type(e).__name__}: {str(e)[:160]}")
            if sp.is_active:
                sp.rollback()
            # objects created inside the failed savepoint are gone -> drop them from the caches
            dies = {k: v for k, v in dies.items() if v in db}
            techs = {k: v for k, v in techs.items() if v in db}
            users_by_name = {k: v for k, v in users_by_name.items() if v in db}
    db.commit()
    return {"created": created, "skipped": skipped, "errors": errors[:30], "total_rows": len(rows)}


def import_dies(db: Session, company, filename: str, data: bytes) -> dict:
    rows = read_rows(filename, data)
    created = updated = 0
    for raw in rows:
        r = {clean(k): v for k, v in raw.items()}
        part_no, process = clean(r.get("PartNo") or r.get("Part No") or r.get("part_no")), clean(r.get("DieProcess") or r.get("Die Process") or r.get("die_process"))
        if not part_no:
            continue
        process = process or "-"                 # stored as "-" when blank, so a re-import finds the same die
        die = db.query(Die).filter(Die.company_id == company.id, Die.part_no == part_no, Die.die_process == process).first()
        vals = dict(die_code=clean(r.get("DieCode") or r.get("Die Code") or r.get("die_code")) or None,
                    customer=clean(r.get("Customer") or r.get("customer")), model=clean(r.get("Model") or r.get("model")),
                    part_name=clean(r.get("PartName") or r.get("Part Name") or r.get("part_name")),
                    line=norm_line(r.get("Line") or r.get("line")),
                    machine_no=clean(r.get("Machine") or r.get("machine_no")))
        for col, key in (("PM Interval Days", "pm_interval_days"), ("PM Interval Strokes", "pm_interval_strokes"), ("Stroke Count", "stroke_count")):
            try:
                v = r.get(col) if r.get(col) not in (None, "") else r.get(key)
                if v not in (None, ""):
                    vals[key] = int(float(str(v).replace(",", "")))
            except (TypeError, ValueError):
                pass
        pl = parse_dt(r.get("PM Last Date") or r.get("pm_last_date"), prefer_us=False)
        if pl:
            vals["pm_last_date"] = pl.date()
        if die:
            for k, v in vals.items():
                if v or (k in ("pm_interval_days", "pm_interval_strokes") and v == 0 and k in vals):
                    setattr(die, k, v)
            updated += 1
        else:
            die = Die(company_id=company.id, part_no=part_no, die_process=process, **vals)
            db.add(die)
            created += 1
        if die.pm_last_date and die.pm_interval_days:
            die.pm_next_date = die.pm_last_date + timedelta(days=die.pm_interval_days)
        if vals["customer"] and not db.query(Customer).filter_by(company_id=company.id, name=vals["customer"]).first():
            db.add(Customer(company_id=company.id, name=vals["customer"]))
            db.flush()
        if vals["line"] and not db.query(ProdLine).filter_by(company_id=company.id, name=vals["line"]).first():
            db.add(ProdLine(company_id=company.id, name=vals["line"]))
            db.flush()
    db.commit()
    return {"created": created, "updated": updated, "total_rows": len(rows)}
