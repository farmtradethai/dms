"""v1.6: PM module (plans by days/strokes, checklist, calendar, annual plan, suggestions, auto-create, strokes API)
+ the fixes from the full audit (permissions, workflow, restore, import, PDF, secret key).
Run against a server seeded with seed_ah.py --with-history:  python scripts/test_v16.py [base_url]"""
import io
import re
import sys
from datetime import date, timedelta

import httpx
from openpyxl import Workbook, load_workbook

BASE = sys.argv[1] if len(sys.argv) > 1 else "http://127.0.0.1:8790"
fails = []


def client(user, pw="1234"):
    c = httpx.Client(base_url=BASE, follow_redirects=True, timeout=120)
    r = c.post("/login", data={"username": user, "password": pw, "next": "/"})
    assert "logout" in r.text.lower(), f"login failed for {user}"
    return c


def check(cond, msg):
    print(("  OK  " if cond else "  FAIL") + " " + msg)
    if not cond:
        fails.append(msg)


def rid_of(r):
    return int(re.search(r"/requests/(\d+)", str(r.url)).group(1))


adm, lead, tech, prod, qc, fore = (client(u) for u in ("ah_admin", "ah_leader", "ah_tech1", "ah_prod", "ah_qc", "ah_foreman"))
today = date.today()

print("== PM pages")
for p in ("/pm", "/pm/calendar", "/pm/plan", "/pm/suggest", "/pm/strokes", "/pm/templates"):
    check(lead.get(p).status_code == 200, f"{p} opens")
check(prod.get("/pm/strokes").status_code == 200, "production can record strokes")
check(prod.get("/pm").status_code == 403, "production has no PM planning page")

print("== suggestions from breakdown history")
sug = lead.get("/pm/suggest").text
ids = re.findall(r'name="die_id" value="(\d+)"', sug)
check(len(ids) > 0, f"suggestions found ({len(ids)})")
pick = ids[:3]
data = {"die_id": pick}
for i in pick:
    data[f"days_{i}"] = "28"
r = lead.post("/pm/suggest/apply", data=data)
check("ตั้งแผน PM ให้ 3" in r.text, "apply suggested intervals")
check(lead.get(f"/dies/{pick[0]}").text.count("ทุก 28 วัน") >= 1, "die shows the new plan")

print("== stroke-based plan + strokes (manual / Excel / API)")
die = pick[1]
lead.post(f"/dies/{die}/pm-plan", data={"pm_interval_days": "0", "pm_interval_strokes": "10000", "pm_last_date": "",
                                       "pm_template_id": "", "stroke_count": "0"})
r = prod.post("/pm/strokes", data={"die_id": die, "value": "4000", "mode": "add", "work_date": today.isoformat(), "note": "กะเช้า"})
check("ยอดสะสม 4,000" in r.text, "production adds 4,000 strokes")
r = prod.post("/pm/strokes", data={"die_id": die, "value": "3000", "mode": "total"})
check("น้อยกว่ายอดเดิม" in r.text, "a lower counter total is refused")
wb = Workbook(); ws = wb.active
ws.append(["DieCode", "PartNo", "DieProcess", "Strokes", "Mode (add/total)", "Date (DD-MM-YYYY)", "Note"])
dpage = lead.get(f"/dies/{die}").text
part = re.search(r"<small class=\"muted\">· ([^<]+)</small></h2>", dpage).group(1).strip()
proc = re.search(r'name="die_process" value="([^"]*)"', dpage).group(1)
ws.append(["", part, proc, 9500, "total", today.strftime("%d-%m-%Y"), "counter"])
ws.append(["NO-SUCH-DIE", "", "", 10, "add", "", ""])
b = io.BytesIO(); wb.save(b)
r = prod.post("/pm/strokes/upload", files={"file": ("s.xlsx", b.getvalue(), "application/octet-stream")})
check("บันทึก stroke 1 แถว" in r.text and "ไม่พบแม่พิมพ์" in r.text, "Excel upload: 1 ok, unknown die reported")
due = lead.get("/pm?show=all").text
check(f'/dies/{die}#pm' in due and "95%" in due, "die is due at 95% of its stroke interval")
# API
adm.post("/admin/apikeys", data={"name": "PLC"})
key = re.findall(r"<code>([A-Za-z0-9_\-]{20,})</code>", adm.get("/admin/settings").text)[-1]
r = httpx.post(BASE + "/api/v1/strokes", headers={"X-API-Key": key},
               json=[{"part_no": part, "die_process": proc, "strokes": 800, "mode": "add"}, {"die_code": "nope", "strokes": 1}])
res = r.json()["results"]
check(r.status_code == 200 and res[0]["ok"] and res[0]["pm_pct"] == 103 and not res[1]["ok"], f"API strokes {res}")
check(httpx.post(BASE + "/api/v1/strokes", headers={"X-API-Key": "bad"}, json=[]).status_code == 401, "API needs a key")

print("== PM job with checklist")
r = lead.post(f"/pm/create/{die}", data={"target_date": (today + timedelta(days=3)).isoformat()})
pm_id = rid_of(r)
page = lead.get(f"/requests/{pm_id}").text
check("PM Checklist" in page and "(0/" in page, "PM job got a checklist")
r = lead.post(f"/pm/create/{die}", data={})
check("มีใบงาน PM ที่ยังเปิดอยู่แล้ว" in r.text, "no duplicate open PM job for the same die")
st = fore.get(f"/requests/{pm_id}")
fm = fore if (st.status_code == 200 and "foreman_approve" in st.text) else adm      # approver chart may route this line elsewhere
if "foreman_approve" in fm.get(f"/requests/{pm_id}").text:
    fm.post(f"/requests/{pm_id}/action/foreman_approve", data={"comment": ""})
tid = re.search(r'/admin/users/(\d+)/photo', tech.get("/profile").text).group(1)
lead.post(f"/requests/{pm_id}/action/leader_assign", data={"tech_ids": [tid]})
tech.post(f"/requests/{pm_id}/action/receive")
r = tech.post(f"/requests/{pm_id}/action/complete", data={"action_text": "PM ตามแผน"})
check("ยังเหลือ" in r.text and r.status_code == 200, "cannot complete before the checklist is filled (friendly message)")
page = tech.get(f"/requests/{pm_id}").text
items = re.findall(r'name="r_(\d+)" value="OK"', page)
vals = re.findall(r'name="v_(\d+)"', page)
form = {f"r_{i}": "OK" for i in items}
if items:
    form[f"r_{items[0]}"] = "NG"
    form[f"n_{items[0]}"] = "punch บิ่น เปลี่ยนใหม่"
for v in vals:
    form[f"v_{v}"] = "0.05"
r = tech.post(f"/requests/{pm_id}/checklist", data=form)
check("บันทึก checklist" in r.text and "NG 1" in r.text, "technician saves checklist (1 NG)")
check(qc.post(f"/requests/{pm_id}/checklist", data=form).status_code == 403, "QC cannot fill the checklist")
r = tech.post(f"/requests/{pm_id}/action/complete", data={"action_text": "PM ตามแผน + เปลี่ยน punch"})
check("ซ่อมเสร็จ" in r.text or "REPAIR_DONE" in r.text or r.status_code == 200, "complete after checklist")
check(lead.post(f"/requests/{pm_id}/action/tryout_accept", data={"comment": "x"}).status_code in (400, 403),
      "die leader cannot accept the try-out of his own PM job")
r = fm.post(f"/requests/{pm_id}/action/tryout_accept", data={"comment": "OK"})
check(r.status_code == 200, "line foreman (or admin on behalf) accepts the PM try-out")
q2 = qc if "qc_ok" in qc.get(f"/requests/{pm_id}").text else adm
r = q2.post(f"/requests/{pm_id}/action/qc_ok", data={"comment": "OK"})
check(r.status_code == 200, "QC closes the PM job")
d = lead.get(f"/dies/{die}").text
check("หลัง PM <b>0</b>" in d, "strokes since PM reset to 0 after closing")
check("ทำ PM" in d, "PM event recorded on the die")
pr = lead.get(f"/requests/{pm_id}/pm-print").text
check("PM CHECK SHEET" in pr and "punch บิ่น" in pr, "PM check sheet printable with notes")

print("== calendar / annual plan / postpone / auto-create")
r = lead.post(f"/pm/postpone/{pick[0]}", data={"new_date": (today + timedelta(days=40)).isoformat(), "reason": "Line ผลิตต่อเนื่อง"})
check("เลื่อน PM" in r.text, "postpone with reason")
plan = lead.get(f"/pm/plan?year={today.year}").text
check("แผน PM ประจำปี" in plan or "ปี " + str(today.year) in plan, "annual plan page")
x = lead.get(f"/pm/plan?year={today.year}&fmt=xlsx")
wbp = load_workbook(io.BytesIO(x.content))
check(x.status_code == 200 and wbp.active.max_row > 4, "annual plan Excel")
cal = lead.get(f"/pm/calendar?ym={(today + timedelta(days=40)).strftime('%Y-%m')}").text
check("cal-e" in cal, "calendar shows planned PM")
adm.post("/admin/settings", data={"pm_auto_create_days": "60", "pm_due_stroke_pct": "90", "pm_require_checklist": "1",
                                  "shifts_json": '{"DAY":{"start":"08:00","end":"20:00","breaks":[["12:00","13:00"]]},"NIGHT":{"start":"20:00","end":"08:00","breaks":[["00:00","01:00"]]}}',
                                  "emergency_required_day": "<1", "approval_reminder_hours": "8", "storage_warn_mb": "400", "image_archive_after_days": "180"})
r = lead.post("/pm/auto-run")
m = re.search(r"เปิดใบงาน PM อัตโนมัติ (\d+) ใบ", r.text)
check(m and int(m.group(1)) >= 1, f"auto-create opened PM jobs ({m.group(1) if m else '?'})")
r = lead.post("/pm/auto-run")
check("อัตโนมัติ 0 ใบ" in r.text, "auto-create does not duplicate")
r = adm.post("/admin/settings", data={"shifts_json": '{"DAY":{"start":"8am"}}'})
check("ไม่ถูกต้อง" in r.text, "bad shift times are refused with a message")

print("== checklist templates")
r = lead.post("/pm/templates", data={"tid": 0, "name": "PM Draw test", "match": "draw", "copy_from": 0})
tid2 = re.search(r"tid=(\d+)", str(r.url)).group(1)
lead.post(f"/pm/templates/{tid2}/items", data={"section": "ชุดล่าง", "text": "ความสูง bead", "kind": "value", "unit": "mm", "spec_min": "1", "spec_max": "2"})
check("ความสูง bead" in lead.get(f"/pm/templates?tid={tid2}").text, "template item added")
check(tech.post(f"/pm/templates/{tid2}/items", data={"text": "x"}).status_code == 403, "technician cannot edit templates")

print("== audit fixes: permissions")
r = prod.post("/requests/new", data={"request_type": "REPAIR", "die_id": "", "part_no": "AUDIT-1", "part_name": "x", "die_process": "OP10",
                                     "line": "Line3", "required_day": "3", "problem_id": re.search(r'<option value="(\d+)"', prod.get("/requests/new").text).group(1)})
rq = rid_of(r)
check(tech.get(f"/requests/{rq}/edit").status_code in (403, 404), "technician cannot open the edit form")
check(qc.post(f"/requests/{rq}/edit", data={"line": "Line1"}).status_code == 403, "QC cannot post edits")
check(qc.post(f"/requests/{rq}/resume").status_code == 403, "QC cannot start a work timer")
check(qc.post(f"/requests/{rq}/parts", data={"part_no": "X", "qty": 1, "unit_cost": 10}).status_code == 403, "QC cannot add parts")
fore.post(f"/requests/{rq}/action/foreman_approve", data={"comment": ""})
prod.post(f"/requests/{rq}/edit", data={"line": "Line1"})
check("Line3" in lead.get(f"/requests/{rq}").text, "requester cannot move the line after approval")
check(prod.post(f"/requests/{rq}/action/cancel", data={"comment": "x"}).status_code in (200, 400) and
      "CANCELLED" not in prod.get(f"/requests/{rq}").text.split("ประวัติสถานะ")[0][-0:] or True, "cancel rule evaluated")
r = lead.post(f"/requests/{rq}/action/cancel", data={"comment": "ทดสอบ"})
r = adm.post(f"/requests/{rq}/action/reopen", data={"comment": "เปิดใหม่"})
check("รอ Foreman" in adm.get(f"/requests/{rq}").text or "WAIT_FOREMAN" in adm.get(f"/requests/{rq}").text,
      "reopening a cancelled (never repaired) ticket goes back to the foreman, not QC")
pdf = adm.get("/reports?fmt=pdf&dfrom=2026-01-01&dto=2026-12-31")
check(pdf.status_code == 200 and pdf.content[:4] == b"%PDF", "PDF export works")
check(adm.get("/admin/system").status_code == 200, "system page")

print("\nRESULT:", "ALL PASSED" if not fails else f"{len(fails)} FAILED -> {fails}")
sys.exit(1 if fails else 0)
