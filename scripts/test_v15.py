"""v1.5: approval routing per company, approver-only visibility, start-up kit + readiness, staff photos, TV data.
Run against a running server seeded with seed_ah.py --with-history:  python scripts/test_v15.py [base_url]"""
import io
import re
import sys
import zipfile

import httpx
from openpyxl import load_workbook
from PIL import Image

BASE = sys.argv[1] if len(sys.argv) > 1 else "http://127.0.0.1:8790"
fails = []
RUN = __import__("random").randint(1000, 9999)


def client(user, pw="1234"):
    c = httpx.Client(base_url=BASE, follow_redirects=True, timeout=120)
    r = c.post("/login", data={"username": user, "password": pw, "next": "/"})
    assert "logout" in r.text.lower(), f"login failed for {user}"
    return c


def check(cond, msg):
    print(("  OK  " if cond else "  FAIL") + " " + msg)
    if not cond:
        fails.append(msg)


def jpg(color=(200, 80, 60), size=(900, 1200)):
    b = io.BytesIO()
    Image.new("RGB", size, color).save(b, "JPEG")
    return b.getvalue()


adm = client("ah_admin")


def make_user(username, name, role, line=""):
    adm.post("/admin/users", data={"uid": 0, "username": username, "full_name": name, "role": role, "password": "Pass#1234",
                                   "line": line, "email": f"{username}@test.local", "is_active": "1"})
    page = adm.get("/admin/approvers").text
    m = re.search(r'<option value="(\d+)">' + re.escape(name), page)
    return int(m.group(1)) if m else None


def uid_of(name):
    m = re.search(r'<option value="(\d+)">' + re.escape(name), adm.get("/admin/approvers").text)
    return int(m.group(1)) if m else None


print("== approval routing (Admin › ผังผู้อนุมัติ)")
check(adm.get("/admin/approvers").status_code == 200, "approver page opens")
fm2 = make_user(f"fm2_{RUN}", f"Foreman สอง {RUN}", "foreman", "Line1")
qc2 = make_user(f"qc2_{RUN}", f"QC สอง {RUN}", "qc", "Line1")
fm_main = uid_of("") if False else None
page = adm.get("/admin/approvers").text
fm_id = int(re.search(r'<option value="(\d+)">[^<]*· Foreman', page).group(1))
fm_user_name = re.search(r'<option value="' + str(fm_id) + r'">([^·<]+)·', page).group(1).strip()
# Line3 foreman step -> the original ah_foreman only ; all other lines -> fm2 ; QC Line3 -> ah_qc only
foreman_me = client("ah_foreman")
me_name = re.search(r"👤 ([^<]+) <small", foreman_me.get("/profile").text).group(1).strip()
ah_fm = uid_of(me_name)
ah_qc_name = re.search(r"👤 ([^<]+) <small", client("ah_qc").get("/profile").text).group(1).strip()
ah_qc = uid_of(ah_qc_name)
for step, lines, users in (("FOREMAN", ["Line3"], [ah_fm]), ("FOREMAN", [""], [fm2]), ("QC", ["Line3"], [ah_qc]), ("QC", [""], [qc2])):
    adm.post("/admin/approvers", data={"step": step, "line": lines, "user_ids": [str(u) for u in users]})
cov = adm.get("/admin/approvers").text
check("กฎทั้งหมด (4)" in cov or "กฎทั้งหมด (" in cov, "rules listed")
check(f"Foreman สอง {RUN}" in cov, "coverage table shows the new approver")

prod, fmA, fmB = client("ah_prod"), foreman_me, client(f"fm2_{RUN}", "Pass#1234")
d = prod.get("/api/dies", params={"q": "44591"}).json()["dies"][0]
pid = re.search(r'<option value="(\d+)"[^>]*>ครีบรู', prod.get("/requests/new").text).group(1)
r = prod.post("/requests/new", data={"request_type": "REPAIR", "die_id": d["id"], "line": "Line3", "required_day": "3", "problem_id": pid})
rid = int(re.search(r"/requests/(\d+)", str(r.url)).group(1))
dr = re.search(r"(AHMD\d+)", r.text).group(1)
check(f'href="/requests/{rid}"' in fmA.get(f"/requests?view=approve&q={dr}").text, "Line3 foreman sees the ticket in 'รอฉันอนุมัติ'")
check(f'href="/requests/{rid}"' not in fmB.get(f"/requests?view=approve&q={dr}").text, "other-line foreman does NOT see it in his queue")
check(fmB.get(f"/requests/{rid}").status_code == 403, "other-line foreman cannot open it (approver sees own lines only)")
r = fmB.post(f"/requests/{rid}/action/foreman_approve", data={"comment": "x"})
check(r.status_code in (400, 403), f"other-line foreman cannot approve ({r.status_code})")
check("foreman_approve" in fmA.get(f"/requests/{rid}").text, "designated foreman has the approve button")
# other line ticket goes to fm2
r = prod.post("/requests/new", data={"request_type": "REPAIR", "die_id": d["id"], "line": "Line1", "required_day": "3", "problem_id": pid})
rid1 = int(re.search(r"/requests/(\d+)", str(r.url)).group(1))
dr1 = re.search(r"(AHMD\d+)", r.text).group(1)
check(f'href="/requests/{rid1}"' in fmB.get(f"/requests?view=approve&q={dr1}").text, "Line1 ticket waits for the all-lines rule person (fm2)")
check(f'href="/requests/{rid1}"' not in fmA.get(f"/requests?view=approve&q={dr1}").text, "Line3 foreman does not get the Line1 ticket")
# manager acts on behalf -> recorded
mgr = client("ah_manager")
r = mgr.post(f"/requests/{rid1}/action/foreman_approve", data={"comment": "ไม่อยู่"})
check(r.status_code == 200 and "อนุมัติแทน" in mgr.get(f"/requests/{rid1}").text, "manager can approve on behalf and it is recorded")
r = fmA.post(f"/requests/{rid}/action/foreman_approve", data={"comment": "ok"})
check(r.status_code == 200, "designated foreman approves")
# QC routing
lead = client("ah_leader")
tech_c = client("ah_tech1")
tid = re.search(r'/admin/users/(\d+)/photo', tech_c.get("/profile").text).group(1)
lead.post(f"/requests/{rid}/action/leader_assign", data={"tech_ids": [tid], "comment": ""})
tech_c.post(f"/requests/{rid}/action/receive")
tech_c.post(f"/requests/{rid}/action/complete", data={"action_text": "เจียร punch"})
r = prod.post(f"/requests/{rid}/action/tryout_accept", data={"comment": "OK"})
check(r.status_code == 200, "requester accepts try-out")
qcB = client(f"qc2_{RUN}", "Pass#1234")
check(qcB.post(f"/requests/{rid}/action/qc_ok", data={"comment": ""}).status_code in (400, 403), "QC of another line cannot judge")
check(f'href="/requests/{rid}"' in client("ah_qc").get(f"/requests?view=approve&q={dr}").text, "Line3 QC sees it waiting")
check(client("ah_qc").post(f"/requests/{rid}/action/qc_ok", data={"comment": "OK"}).status_code == 200, "Line3 QC closes it")
check("Foreman / Engineer อนุมัติใบแจ้ง" in fmB.get("/profile").text, "profile tells the approver what he approves")
# visibility switch
adm.post("/admin/approvers/visibility", data={"vis": "all"})
check(fmB.get(f"/requests/{rid}").status_code == 200, "visibility 'all' lets approvers browse every ticket")
check(f'href="/requests/{rid}"' not in fmB.get(f"/requests?view=approve&q={dr}").text, "...but 'รอฉันอนุมัติ' stays personal")
adm.post("/admin/approvers/visibility", data={"vis": "own"})

print("== start-up kit for a NEW company")
sa = client("admin", "admin1234")
code = f"N{RUN}"
r = sa.post("/admin/companies", data={"cid": 0, "code": code, "name": f"New Co {RUN}", "dr_prefix": "NC", "dr_style": "running",
                                       "logo_text": "AAPICO", "is_active": "1", "masters": "empty",
                                       "admin_username": f"adm_{RUN}", "admin_name": "New Admin", "admin_password": "Start#2026"})
check("/admin/setup" in str(r.url), "after creating a company the superadmin lands on its start-up page")
na = httpx.Client(base_url=BASE, follow_redirects=True, timeout=120)
na.post("/login", data={"username": f"adm_{RUN}", "password": "Start#2026", "next": "/"})
na.post("/profile", data={"lang": "th", "theme": "light", "accent": "#1f5fbf", "password": "NewPass#99", "password2": "NewPass#99"})
page = na.get("/admin/setup").text
check("ยังขาดข้อมูลจำเป็น" in page, "empty company is reported as NOT ready")
kit = na.get("/admin/setup/kit.xlsx").content
wb = load_workbook(io.BytesIO(kit))
need = ["0_อ่านก่อน", "1_Company", "2_Lines", "3_Customers", "4_Problems", "5_RootCauses", "6_Users", "7_Approvers", "8_Dies", "9_History"]
check(all(s in wb.sheetnames for s in need), f"template has every sheet {wb.sheetnames}")
# make usernames unique for this run
ws = wb["6_Users"]
for row in ws.iter_rows(min_row=2):
    row[0].value = f"{row[0].value}_{RUN}"
for row in wb["7_Approvers"].iter_rows(min_row=2):
    row[2].value = f"{row[2].value}_{RUN}"
bad = io.BytesIO()
wb["8_Dies"].append(["X", "M", "", "no part", "OP10", "", "Line3", "", "abc", "99-99-2026", "Y"])     # 2 errors
wb.save(bad)
r = na.post("/admin/setup/upload", files={"file": ("kit.xlsx", bad.getvalue(), "application/octet-stream")})
check("พบข้อผิดพลาด" in r.text and "8_Dies" in r.text and "PartNo" in r.text, "validation lists row errors")
check("ยังขาดข้อมูลจำเป็น" in na.get("/admin/setup").text, "...and nothing was saved")
wb["8_Dies"].delete_rows(wb["8_Dies"].max_row)
good = io.BytesIO()
wb.save(good)
r = na.post("/admin/setup/upload", files={"file": ("kit.xlsx", good.getvalue(), "application/octet-stream")})
check("ไฟล์ถูกต้องทั้งหมด" in r.text, "corrected file validates")
tok = re.search(r'name="token" value="([^"]+)"', r.text).group(1)
r = na.post("/admin/setup/upload", data={"token": tok, "do_import": "1"})
check("นำเข้าเรียบร้อย" in r.text, "import done")
check("พร้อมใช้งานจริง" in r.text and "ยังขาดข้อมูลจำเป็น" not in r.text, "company is READY after the kit import")
cur = load_workbook(io.BytesIO(na.get("/admin/setup/kit.xlsx?data=1").content))
check(cur["8_Dies"].max_row == 2 and cur["6_Users"].max_row >= 7, "kit with current data round-trips")
r = na.post("/admin/setup/upload", files={"file": ("kit.xlsx", good.getvalue(), "application/octet-stream")})
tok = re.search(r'name="token" value="([^"]+)"', r.text).group(1)
r = na.post("/admin/setup/upload", data={"token": tok, "do_import": "1"})
check("+0" in na.get("/admin/audit?action=setup_kit_import").text or "นำเข้าเรียบร้อย" in r.text, "re-import updates instead of duplicating")
check(len(re.findall(r"<tr", na.get("/dies").text)) <= 4, "no duplicate dies after re-import")
check(dr not in na.get("/requests?view=all").text, "new company cannot see AH tickets")
check(code not in adm.get("/admin/approvers").text, "AH admin page untouched by the new company")
check(na.get(f"/requests/{rid}").status_code in (403, 404), "new company cannot open an AH ticket")

print("== staff photos")
me = client("ah_tech1")
my_id = int(re.search(r'/admin/users/(\d+)/photo', me.get("/profile").text).group(1))
r = me.post(f"/admin/users/{my_id}/photo", files={"file": ("me.jpg", jpg(), "image/jpeg")}, data={"back": "/profile"})
check("บันทึกรูป" in r.text and 'class="av"' in r.text, "technician uploads own photo from profile")
other = re.search(r'/admin/users/(\d+)/photo', adm.get("/admin/users").text).group(1)
check(me.post(f"/admin/users/{other}/photo", files={"file": ("x.jpg", jpg(), "image/jpeg")}).status_code in (403,) or other == str(my_id),
      "cannot change someone else's photo")
tv = adm.get("/api/tv").json()
check(any(p["photo"] for p in tv["team"]), "TV team shows the photo")
src = re.search(r'<img class="av" src="([^"]+)"', me.get("/profile").text).group(1)
img = adm.get(src)
check(img.status_code == 200 and Image.open(io.BytesIO(img.content)).size == (320, 320), "photo stored as 320x320")
check(httpx.get(BASE + src).status_code in (401, 403) or "login" in str(httpx.get(BASE + src, follow_redirects=True).url), "photo needs login")
z = io.BytesIO()
with zipfile.ZipFile(z, "w") as f:
    f.writestr("ah_tech2.jpg", jpg((20, 120, 200)))
    f.writestr("unknown_person.jpg", jpg())
r = adm.post("/admin/users/photos-zip", files={"file": ("p.zip", z.getvalue(), "application/zip")})
check("บันทึกรูป 1 คน" in r.text and "unknown_person.jpg" in r.text, "ZIP upload matches by username, reports the rest")

print("== TV board data")
check(set(tv["kpi"]) >= {"open", "emergency", "overdue", "approve", "tryout", "qc", "mttr_h", "ontime"}, "TV KPIs")
check(all({"techs", "wait", "status_label"} <= set(j) for j in tv["jobs"][:5]), "TV jobs carry technicians + who it waits for")
check(adm.get("/tv").status_code == 200 and "Scroller" in adm.get("/tv").text, "TV page has scrolling controls")

print("== v1.5.1 dates DD-MM-YYYY + compact cards")
js = adm.get("/static/js/app.js").text
check("dmyDates" in js and "DD-MM-YYYY" in js, "date fields shown as DD-MM-YYYY")
check("?v=" in adm.get("/").text, "css/js versioned so browsers load the new style after an upgrade")
xl = load_workbook(io.BytesIO(adm.get("/analytics?fmt=xlsx").content))
check(all(not re.match(r"^\d{4}-\d{2}", str(c.value or "")) for row in xl["Lead time ต่อขั้น"].iter_rows() for c in row), "Excel months as MM-YYYY")

print("\nRESULT:", "ALL PASSED" if not fails else f"{len(fails)} FAILED -> {fails}")
sys.exit(1 if fails else 0)
